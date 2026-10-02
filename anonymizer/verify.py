"""Проверка готового файла: открывается ли он и не сломаны ли ссылки между листами.

Файл, который не открылся у самой программы, не откроется и у пользователя. Лучше сообщить об этом сразу,
чем отдать повреждённый результат.
"""

from __future__ import annotations

import re
from pathlib import Path

SHEET_REF = re.compile(r"(?:^|[^\w\]'])(?:'((?:[^']|'')+)'|([A-Za-zА-Яа-яЁё_][\w.]*))!")


def check_openable(path: Path) -> tuple[bool, str]:
    """(файл открывается, пояснение). Пояснение пустое, если всё в порядке."""
    suffix = path.suffix.lower()
    try:
        if suffix == ".xlsx":
            return _check_xlsx(path)
        if suffix == ".docx":
            from docx import Document
            Document(str(path))
        elif suffix == ".pptx":
            from pptx import Presentation
            Presentation(str(path))
        elif suffix == ".pdf":
            import fitz
            fitz.open(path).close()
    except Exception as exc:
        return False, f"Файл не открывается стандартными средствами ({type(exc).__name__}: {str(exc)[:120]})"
    return True, ""


def _check_xlsx(path: Path) -> tuple[bool, str]:
    from openpyxl import load_workbook
    workbook = load_workbook(path)
    names = {name.casefold() for name in workbook.sheetnames}
    missing: set[str] = set()
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("=") and "!" in value:
                    for match in SHEET_REF.finditer(value):
                        target = (match.group(1) or match.group(2) or "").replace("''", "'")
                        if target and target.casefold() not in names and not target.startswith("["):
                            missing.add(target)
    if missing:
        listed = ", ".join(sorted(missing)[:5])
        return True, f"Формулы ссылаются на листы, которых нет в книге: {listed}"
    return True, ""
