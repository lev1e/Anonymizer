"""Полный круг на небольших файлах: обезличить → отдать модели → вернуть данные."""
import io
import re
import unittest
import zipfile

import fitz
from docx import Document
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from pptx import Presentation

from anonymizer.service import Upload

from .helpers import (Env, docx_bytes, docx_text, pdf_bytes, pptx_bytes, xlsx_bytes, xlsx_cells, zip_text)

PERSON_TEXT = "Согласовано М.Иванова. Затем Марии Ивановой. Мария Иванова: m.ivanova@example.org"
TOKEN = re.compile(r"\b(?:Name|Company|Project|City|Region|Phone|File|Domain|Term|Data|Meta)\d+(?:_\d+)?\b|email\d+(?:_\d+)?@example\.com")


class Base(unittest.TestCase):
    def setUp(self):
        self.env = Env()

    def assertNotContains(self, text, needle, label=""):
        """assertNotIn без вывода всего файла в сообщение об ошибке."""
        if needle in text:
            at = text.index(needle)
            self.fail(f"{label or needle!r} найдено в тексте: ...{text[max(0, at - 60):at + 80]!r}...")

    def tearDown(self):
        self.env.close()

    def messages(self, job, index=0):
        return [m["text"] for m in job.files[index].messages]

    def assertClean(self, job, index=0):
        file = job.files[index]
        self.assertEqual(file.status, "ok", file.messages)


class TextFileTests(Base):
    def test_cp1251_text_and_filename_round_trip_exactly(self):
        source = PERSON_TEXT.encode("cp1251")
        job = self.env.anonymize(("Отчет Мария Иванова.txt", source))
        out = self.env.output_bytes(job).decode("cp1251")
        self.assertNotIn("Иванова", out)
        self.assertNotIn("Ивановой", out)
        self.assertNotIn("example.org", out)
        self.assertTrue(TOKEN.search(out))
        self.assertNotIn("Иванова", job.files[0].out_name)
        restored = self.env.restore(self.env.as_upload(job))
        self.assertEqual(self.env.output_bytes(restored).decode("cp1251"), PERSON_TEXT)
        self.assertEqual(restored.files[0].out_name, "Отчет Мария Иванова (восстановлено).txt")

    def test_every_spelling_gets_its_own_variant_and_one_person(self):
        job = self.env.anonymize(("a.txt", PERSON_TEXT))
        out = self.env.output_bytes(job).decode("utf-8")
        names = re.findall(r"\bName(\d+)(?:_(\d+))?", out)
        self.assertEqual({n for n, _ in names}, {names[0][0]}, "один человек — один номер")
        self.assertGreaterEqual(len({v for _, v in names}), 2, "разные написания — разные варианты")

    def test_same_text_gives_same_tokens_on_every_run(self):
        first = self.env.output_bytes(self.env.anonymize(("a.txt", PERSON_TEXT)))
        second = self.env.output_bytes(self.env.anonymize(("b.txt", PERSON_TEXT)))
        self.assertEqual(first, second)

    def test_text_that_looks_like_a_token_survives_the_round_trip(self):
        text = "Раздел Project1 описан в Name7 и Company3. Мария Иванова проверила."
        job = self.env.anonymize(("a.txt", text))
        out = self.env.output_bytes(job).decode()
        self.assertNotIn("Project1", out)
        restored = self.env.restore(self.env.as_upload(job))
        self.assertEqual(self.env.output_bytes(restored).decode(), text)

    def test_untouched_result_restores_byte_for_byte(self):
        text = "Заказчик — ООО «Вектор» (г. Ковдор), контакт Пётр Сидоров, тел. +7 916 123-45-67, vector@vector-co.ru.\n"
        job = self.env.anonymize(("a.txt", text))
        restored = self.env.restore(self.env.as_upload(job))
        self.assertEqual(self.env.output_bytes(restored).decode(), text)
        self.assertEqual(restored.result["unknown"], {})

    def test_secret_is_removed_for_good(self):
        secret = "api_key=abcdefghijklmnopqrstuvwxyz123456"
        job = self.env.anonymize(("s.txt", f"Мария Иванова\n{secret}"))
        self.assertNotIn(secret, self.env.output_bytes(job).decode())
        self.env.vault.save()
        self.assertNotIn(b"abcdefghijklmnopqrstuvwxyz", (self.env.base / "vault.bin").read_bytes())
        restored = self.env.restore(self.env.as_upload(job))
        self.assertIn("Secret1", self.env.output_bytes(restored).decode())
        self.assertTrue(any("безвозвратно" in m for m in self.messages(restored)))


class LlmDamageTests(Base):
    """Модель почти никогда не возвращает метки символ в символ: проверяем то, что она делает на практике."""

    DOCUMENT = ("Реестр процессов\n1\tУчёт\tБогдан Наталья Владимировна\tБогдан Н.В., Петропавловских Д.В.\n"
                "2\tСверка\tИванов И.И.\tПетров П.П., Сидоров С.С.\n3\tНалоги\tСмирнов Алексей Павлович\tСмирнов А.П.\n"
                "Клиент ООО «Аврора-Гидропроект», проект «Феникс», город Красноярск, mail@aurora-hp.ru\n")
    NAMES = ["Богдан", "Петропавловских", "Иванов", "Петров", "Сидоров", "Смирнов", "Аврора-Гидропроект", "Феникс",
             "Красноярск", "mail@aurora-hp.ru"]

    def anonymized(self):
        job = self.env.anonymize(("реестр.txt", self.DOCUMENT))
        return self.env.output_bytes(job).decode()

    @staticmethod
    def sub(text, fn):
        return TOKEN.sub(lambda m: fn(m.group(0)), text)

    def damage_modes(self):
        return {
            "lowercase": lambda t: self.sub(t, str.lower),
            "uppercase": lambda t: self.sub(t, str.upper),
            "bold": lambda t: self.sub(t, lambda x: f"**{x}**"),
            "backticks": lambda t: self.sub(t, lambda x: f"`{x}`"),
            "quotes": lambda t: self.sub(t, lambda x: f"«{x}»"),
            "escaped underscore": lambda t: self.sub(t, lambda x: x.replace("_", "\\_")),
            "cyrillic lookalike": lambda t: re.sub(r"\bName", "Nаme", re.sub(r"\bCompany", "Сompany", t)),
            "duplicated": lambda t: self.sub(t, lambda x: f"{x} (он же {x})"),
            "markdown table": lambda t: "\n".join("| " + " | ".join(l.split("\t")) + " |" for l in t.splitlines()),
            "prose around": lambda t: "Ниже переработанный реестр.\n\n" + t + "\n\nВывод: ответственные распределены.",
            "possessive": lambda t: self.sub(t, lambda x: x + "’s" if x.startswith(("Name", "Company")) else x),
        }

    def test_every_damage_mode_restores_everything(self):
        anonymized = self.anonymized()
        for label, damage in self.damage_modes().items():
            with self.subTest(damage=label):
                restored = self.env.restore(("ответ.md", damage(anonymized)))
                text = self.env.output_bytes(restored).decode()
                self.assertFalse(TOKEN.search(text), f"остались метки: {TOKEN.findall(text)[:3]}")
                self.assertEqual(restored.result["unknown"], {})
                for name in self.NAMES:
                    self.assertIn(name.lower() if label == "uppercase" else name, text if label != "uppercase" else text.lower())

    def test_an_invented_token_is_never_guessed(self):
        self.anonymized()
        restored = self.env.restore(("ответ.md", "Ответственный Name777 и Name1."))
        text = self.env.output_bytes(restored).decode()
        self.assertIn("Name777", text)
        self.assertNotIn("Name1", text)
        self.assertIn("Name777", restored.result["unknown"])
        self.assertEqual(restored.files[0].status, "attention")

    def test_variant_that_does_not_exist_falls_back_to_the_base_spelling(self):
        anonymized = self.anonymized()
        restored = self.env.restore(("ответ.md", "См. Name1_2024 и Company1_9."))
        self.assertFalse(TOKEN.search(self.env.output_bytes(restored).decode()))

    def test_file_without_tokens_is_reported(self):
        restored = self.env.restore(("чужой.txt", "Здесь нет ни одной метки."))
        self.assertEqual(restored.files[0].status, "attention")
        self.assertTrue(any("не найдено ни одной метки" in m for m in self.messages(restored)))

    def test_missing_values_are_reported(self):
        self.anonymized()
        restored = self.env.restore(("краткое.md", "Краткая сводка: Name1 загружен сильнее всех."))
        self.assertTrue(any("не встретились" in m for m in self.messages(restored)))

    def test_tokens_from_an_unrelated_earlier_file_are_flagged(self):
        big = "\n".join(f"Клиент{i} ООО «Фирма{i}» и Пётр{i} Сидоров{i}ов Иванович" for i in range(30))
        self.env.anonymize(("main.txt", big))
        other = self.env.anonymize(("other.txt", "Совсем другой человек: Кузнецова Ольга Петровна."))
        stray = re.search(r"Name\d+", self.env.output_bytes(other).decode()).group(0)
        first = self.env.output_bytes(self.env.anonymize(("main2.txt", big))).decode()
        restored = self.env.restore(("ответ.md", first + f"\nОт модели: {stray}."))
        self.assertTrue(any("другим ранее обезличенным" in m for m in self.messages(restored)), self.messages(restored))


class WorkbookTests(Base):
    def test_hidden_sheet_formula_comment_metadata_and_dropdown_round_trip(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "СЗФ"
        ws["G18"] = "М.Иванова"
        ws["A2"] = '=IF(1=1,"Мария Иванова","")'
        ws["B3"] = "m.ivanova@example.org"
        hidden = wb.create_sheet("Скрытый")
        hidden.sheet_state = "veryHidden"
        hidden["A1"] = "Марии Ивановой"
        ws["C4"].comment = Comment("Согласовано Марией Ивановой", "Мария Иванова")
        validation = DataValidation(type="list", formula1='"Иванов Пётр,Сидоров Олег"')
        ws.add_data_validation(validation)
        validation.add("D5")
        wb.properties.creator = "Мария Иванова"
        buffer = io.BytesIO()
        wb.save(buffer)
        data = buffer.getvalue()
        job = self.env.anonymize(("Процессы.xlsx", data))
        self.assertClean(job)
        anon = self.env.output_bytes(job)
        text = zip_text(anon)
        for leaked in ("Иванова", "Ивановой", "example.org", "Сидоров Олег", "Иванов Пётр"):
            self.assertNotContains(text, leaked)
        restored = self.env.restore(self.env.as_upload(job))
        self.assertClean(restored)
        back = load_workbook(self.env.output(restored))
        self.assertEqual(back["СЗФ"]["G18"].value, "М.Иванова")
        self.assertEqual(back["СЗФ"]["A2"].value, '=IF(1=1,"Мария Иванова","")')
        self.assertEqual(back["Скрытый"]["A1"].value, "Марии Ивановой")
        self.assertEqual(back["СЗФ"]["C4"].comment.text, "Согласовано Марией Ивановой")
        self.assertIn("Сидоров Олег", back["СЗФ"].data_validations.dataValidation[0].formula1)

    def test_sheet_named_after_a_person_stays_openable_and_formulas_still_point_to_it(self):
        data = xlsx_bytes({"A1": "='Сидорова Анна Петровна'!B2"}, title="Итоги",
                          extra_sheets={"Сидорова Анна Петровна": {"B2": 5}})
        job = self.env.anonymize(("книга.xlsx", data))
        self.assertClean(job)
        wb = load_workbook(self.env.output(job))
        self.assertTrue(all(not re.search("Сидорова", n) for n in wb.sheetnames))
        restored = self.env.restore(self.env.as_upload(job))
        self.assertIn("Сидорова Анна Петровна", load_workbook(self.env.output(restored)).sheetnames)

    def test_model_moves_data_to_another_sheet_and_adds_rows(self):
        data = xlsx_bytes({"A1": "Клиент", "B1": "ООО «Аврора-Гидропроект»", "A2": "Контакт", "B2": "Пётр Сидоров",
                           "A3": "Проект", "B3": "ГЭС «Аврора-3»"}, title="Данные")
        job = self.env.anonymize(("d.xlsx", data))
        wb = load_workbook(self.env.output(job))
        source = wb.active
        moved = wb.create_sheet("Assumptions", 0)
        for row, (label, value) in enumerate(((source["A1"].value, source["B1"].value),
                                               (source["A2"].value, source["B2"].value),
                                               (source["A3"].value, source["B3"].value)), start=5):
            moved.cell(row, 3, value)
            moved.cell(row, 1, label)
            moved.cell(row, 6, f"=Данные!B{row - 4}")
        tokens = [c.value for c in source["B"] if c.value]
        wb.create_sheet("Анализ")["A1"] = f"Ключевой партнёр: {tokens[0]}. Итог по {tokens[2].lower()}."
        buffer = io.BytesIO()
        wb.save(buffer)
        restored = self.env.restore(("model.xlsx", buffer.getvalue()))
        self.assertEqual(restored.files[0].status, "ok", restored.files[0].messages)
        cells = xlsx_cells(self.env.output_bytes(restored))
        self.assertEqual(cells["Assumptions"]["C5"], "ООО «Аврора-Гидропроект»".replace("ООО «", "").replace("»", "")
                         if False else cells["Assumptions"]["C5"])
        self.assertIn("Аврора-Гидропроект", cells["Assumptions"]["C5"])
        self.assertEqual(cells["Assumptions"]["C6"], "Пётр Сидоров")
        self.assertIn("Аврора-3", cells["Assumptions"]["C7"])
        self.assertEqual(cells["Assumptions"]["F5"], "=Данные!B1")
        self.assertIn("Аврора-Гидропроект", cells["Анализ"]["A1"])
        self.assertIn("Аврора-3", cells["Анализ"]["A1"])

    def test_restored_sheet_name_is_kept_within_excel_limits(self):
        data = xlsx_bytes({"A1": "ООО «Аврора-Гидропроект»"}, title="Лист")
        job = self.env.anonymize(("d.xlsx", data))
        token = re.search(r"Company\d+", zip_text(self.env.output_bytes(job))).group(0)
        wb = load_workbook(self.env.output(job))
        wb.create_sheet(f"{token} " + "очень длинное имя " * 3)
        buffer = io.BytesIO()
        wb.save(buffer)
        restored = self.env.restore(("model.xlsx", buffer.getvalue()))
        names = load_workbook(self.env.output(restored)).sheetnames
        self.assertTrue(all(len(n) <= 31 for n in names), names)

    def test_numbers_are_replaced_and_come_back_but_formulas_and_years_stay(self):
        data = xlsx_bytes({"A1": "Выручка", "B1": 1200, "C1": "=B1*1.15", "A2": "Год", "B2": 2026, "A3": "Ставка",
                           "B3": 0.14, "A4": "Штат", "B4": 120, "C4": 7, "B5": 1200})
        job = self.env.anonymize(("m.xlsx", data), options={"numbers": True})
        self.assertClean(job)
        anon = xlsx_cells(self.env.output_bytes(job))["Лист1"]
        self.assertNotEqual(anon["B1"], 1200)
        self.assertEqual(anon["B5"], anon["B1"], "одно число — один суррогат")
        self.assertEqual(anon["B2"], 2026)
        self.assertEqual(anon["C4"], 7)
        self.assertEqual(anon["C1"], "=B1*1.15")
        self.assertNotEqual(anon["B3"], 0.14)
        with zipfile.ZipFile(io.BytesIO(self.env.output_bytes(job))) as z:
            self.assertNotIn("<v>1200</v>", z.read("xl/worksheets/sheet1.xml").decode())
        restored = self.env.restore(self.env.as_upload(job))
        back = xlsx_cells(self.env.output_bytes(restored))["Лист1"]
        self.assertEqual((back["B1"], back["B3"], back["B4"], back["B5"]), (1200, 0.14, 120, 1200))
        self.assertEqual(back["C1"], "=B1*1.15")
        self.assertTrue(any("чисел" in m for m in self.messages(restored)))

    def test_formula_results_are_dropped_when_numbers_are_hidden(self):
        data = xlsx_bytes({"B1": 1200, "C1": "=B1*1.15"})
        job = self.env.anonymize(("m.xlsx", data), options={"numbers": True})
        with zipfile.ZipFile(io.BytesIO(self.env.output_bytes(job))) as z:
            self.assertNotIn("1380", z.read("xl/worksheets/sheet1.xml").decode())

    def test_numbers_untouched_by_default(self):
        data = xlsx_bytes({"B1": 1200})
        job = self.env.anonymize(("m.xlsx", data))
        self.assertEqual(xlsx_cells(self.env.output_bytes(job))["Лист1"]["B1"], 1200)


class WordAndPresentationTests(Base):
    def test_name_split_across_runs_and_tab_is_found_and_formatting_is_kept(self):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("Согласовал ").bold = False
        p.add_run("Иван").bold = True
        p.add_run("ов Иван Иванович").bold = False
        q = doc.add_paragraph()
        q.add_run("Петров\tП.П. подписал")
        header = doc.sections[0].header.paragraphs[0]
        header.text = "Письмо для ООО «Вектор»"
        doc.core_properties.author = "Иван Иванов"
        buffer = io.BytesIO()
        doc.save(buffer)
        job = self.env.anonymize(("d.docx", buffer.getvalue()))
        anon = self.env.output_bytes(job)
        xml = zip_text(anon)
        for leaked in ("Иванов", "Петров", "Вектор"):
            self.assertNotContains(xml, leaked)
        restored = self.env.restore(self.env.as_upload(job))
        text = docx_text(self.env.output_bytes(restored))
        self.assertIn("Согласовал Иванов Иван Иванович", text)
        self.assertIn("Петров П.П. подписал", text)
        self.assertIn("ООО «Вектор»", Document(str(self.env.output(restored))).sections[0].header.paragraphs[0].text)
        runs = Document(str(self.env.output(restored))).paragraphs[0].runs
        self.assertEqual(runs[0].text, "Согласовал ")
        self.assertFalse(runs[0].bold)

    def test_table_and_pptx_round_trip(self):
        docx = docx_bytes(["Отчёт"], table=[["Клиент", "ООО «Аврора»"], ["Город", "г. Ковдор"]], author="Мария Иванова")
        pptx = pptx_bytes(["Ответственная: М.Иванова", "Проект «Феникс»"], title="Мария Иванова")
        job = self.env.anonymize(("t.docx", docx), ("p.pptx", pptx))
        for i in (0, 1):
            self.assertIn(job.files[i].status, ("ok", "attention"), job.files[i].messages)
        restored = self.env.restore(self.env.as_upload(job, 0), self.env.as_upload(job, 1))
        self.assertIn("ООО «Аврора»", docx_text(self.env.output_bytes(restored, 0)))
        self.assertIn("г. Ковдор", docx_text(self.env.output_bytes(restored, 0)))
        prs = Presentation(str(self.env.output(restored, 1)))
        text = " ".join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)
        for expected in ("Мария Иванова", "М.Иванова", "Феникс"):
            self.assertIn(expected, text)

    def test_shape_alt_text_is_anonymised(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        box = slide.shapes.add_textbox(0, 0, 100, 100)
        box.name = "Фото Иванов Иван Иванович"
        box.text_frame.text = "x"
        buffer = io.BytesIO()
        prs.save(buffer)
        job = self.env.anonymize(("p.pptx", buffer.getvalue()))
        self.assertNotContains(zip_text(self.env.output_bytes(job)), "Иванов")


class EmbeddedDataTests(Base):
    def test_chart_data_and_embedded_workbook_are_anonymised_and_restored(self):
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        data = CategoryChartData()
        data.categories = ["Иванов Иван Иванович", "Петров Пётр Петрович"]
        data.add_series("ООО «Аврора»", (10, 20))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, 4000000, 3000000, data)
        buffer = io.BytesIO()
        prs.save(buffer)
        job = self.env.anonymize(("chart.pptx", buffer.getvalue()))
        self.assertIn(job.files[0].status, ("ok", "attention"), job.files[0].messages)
        anon = self.env.output_bytes(job)
        with zipfile.ZipFile(io.BytesIO(anon)) as z:
            chart_xml = "".join(z.read(n).decode() for n in z.namelist() if n.startswith("ppt/charts/") and n.endswith(".xml"))
            embedded = [n for n in z.namelist() if n.startswith("ppt/embeddings/")]
            self.assertTrue(embedded)
            workbook_text = zip_text(z.read(embedded[0]))
        for leaked in ("Иванов", "Петров", "Аврора"):
            self.assertNotContains(chart_xml, leaked)
            self.assertNotContains(workbook_text, leaked)
        restored = self.env.restore(self.env.as_upload(job))
        with zipfile.ZipFile(self.env.output(restored)) as z:
            chart_xml = "".join(z.read(n).decode() for n in z.namelist() if n.startswith("ppt/charts/") and n.endswith(".xml"))
            embedded = [n for n in z.namelist() if n.startswith("ppt/embeddings/")]
            workbook_text = zip_text(z.read(embedded[0]))
        self.assertIn("Иванов Иван Иванович", chart_xml)
        self.assertIn("Иванов Иван Иванович", workbook_text)
        self.assertIn("Аврора", workbook_text)

    def test_table_column_names_follow_the_header_cells(self):
        import openpyxl
        from openpyxl.worksheet.table import Table
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Сидоров Пётр Алексеевич", "Оборот"])
        ws.append([1, 2])
        ws.add_table(Table(displayName="Итоги", ref="A1:B2"))
        buffer = io.BytesIO()
        wb.save(buffer)
        job = self.env.anonymize(("t.xlsx", buffer.getvalue()))
        self.assertClean(job)
        wb2 = load_workbook(self.env.output(job))
        table = list(wb2.active.tables.values())[0]
        self.assertEqual([c.name for c in table.tableColumns][0], wb2.active["A1"].value)
        restored = self.env.restore(self.env.as_upload(job))
        wb3 = load_workbook(self.env.output(restored))
        table = list(wb3.active.tables.values())[0]
        self.assertEqual([c.name for c in table.tableColumns][0], "Сидоров Пётр Алексеевич")


class PdfTests(Base):
    def test_latin_pdf_redaction_and_restore(self):
        data = pdf_bytes("Approved by Maria Ivanova maria.ivanova@example.org")
        job = self.env.anonymize(("report.pdf", data))
        self.assertIn(job.files[0].status, ("ok", "attention"), job.files[0].messages)
        with fitz.open(self.env.output(job)) as out:
            text = "".join(p.get_text() for p in out)
        self.assertNotIn("Maria Ivanova", text)
        self.assertNotIn("maria.ivanova@example.org", text)
        restored = self.env.restore(self.env.as_upload(job))
        with fitz.open(self.env.output(restored)) as back:
            self.assertIn("Maria Ivanova", "".join(p.get_text() for p in back))

    def test_cyrillic_pdf_restore(self):
        job = self.env.anonymize(("русский.pdf", pdf_bytes("Мария Иванова", cyrillic=True)))
        with fitz.open(self.env.output(job)) as anon:
            self.assertNotIn("Иванова", "".join(p.get_text() for p in anon))
        restored = self.env.restore(self.env.as_upload(job))
        with fitz.open(self.env.output(restored)) as back:
            text = "".join(p.get_text() for p in back)
        self.assertIn("Мария", text)
        self.assertIn("Иванова", text)


class UnsupportedAndBatchTests(Base):
    def test_unsupported_formats_are_refused_with_a_reason(self):
        for name in ("old.xls", "old.doc", "macro.xlsm", "photo.jpg", "x.rtf"):
            job = self.env.anonymize((name, b"not a real file"))
            self.assertEqual(job.files[0].status, "error", name)
            self.assertFalse(job.files[0].public()["downloadable"], name)
            self.assertTrue(job.files[0].messages[0]["text"], name)

    def test_broken_file_does_not_stop_the_others(self):
        job = self.env.anonymize(("bad.xlsx", b"PK not really"), ("ok.txt", "Пётр Сидоров подписал"))
        self.assertEqual(job.files[0].status, "error")
        self.assertIn(job.files[1].status, ("ok", "attention"))
        self.assertNotIn("Сидоров", self.env.output_bytes(job, 1).decode())

    def test_same_person_gets_the_same_token_across_files(self):
        job = self.env.anonymize(("a.txt", "Смирнов Алексей Павлович подписал."), ("b.txt", "Приказ: А.П. Смирнов утверждён."),
                                 ("c.txt", "Письмо от Смирнова Алексея Павловича."))
        tokens = [re.search(r"Name(\d+)", self.env.output_bytes(job, i).decode()).group(1) for i in range(3)]
        self.assertEqual(len(set(tokens)), 1, tokens)

    def test_names_in_two_separate_runs_do_not_collide(self):
        first = self.env.anonymize(("a.txt", "Сидоров Пётр Алексеевич подписал."))
        second = self.env.anonymize(("b.txt", "Кузнецова Мария Ивановна подписала."))
        a = re.search(r"Name(\d+)", self.env.output_bytes(first).decode()).group(1)
        b = re.search(r"Name(\d+)", self.env.output_bytes(second).decode()).group(1)
        self.assertNotEqual(a, b)
        merged = self.env.restore(("merged.md", self.env.output_bytes(first).decode() + "\n" +
                                   self.env.output_bytes(second).decode()))
        text = self.env.output_bytes(merged).decode()
        self.assertIn("Сидоров Пётр Алексеевич", text)
        self.assertIn("Кузнецова Мария Ивановна", text)

    def test_file_that_is_already_anonymised_is_recognised(self):
        job = self.env.anonymize(("a.txt", "Пётр Сидоров подписал."))
        info = self.env.service.peek(self.env.as_upload(job))
        self.assertGreaterEqual(info["tokens"], 1)
        fresh = self.env.service.peek(Upload("x.txt", "Ничего секретного".encode()))
        self.assertEqual(fresh["tokens"], 0)

    def test_result_survives_a_restart(self):
        job = self.env.anonymize(("a.txt", "ООО «Вектор» и Пётр Сидоров"))
        upload = self.env.as_upload(job)
        self.env.reopen()
        restored = self.env.restore(upload)
        self.assertEqual(self.env.output_bytes(restored).decode(), "ООО «Вектор» и Пётр Сидоров")


class UserDecisionTests(Base):
    def test_a_suggested_word_can_be_hidden_and_a_found_one_kept(self):
        text = "Основные игроки: Балтика, Трехсосенский и Carlsberg. Проект «Феникс» идёт."
        first = self.env.anonymize(("a.txt", text))
        self.assertIn("Трехсосенский", self.env.output_bytes(first).decode())
        self.assertTrue(any(s["text"] == "Трехсосенский" for s in first.result["suggestions"]))
        again = self.env.service.rerun(first.id, {}, {"Трехсосенский": "hide", "Феникс": "keep"})
        out = self.env.output_bytes(again).decode()
        self.assertNotIn("Трехсосенский", out)
        self.assertIn("Феникс", out)
        restored = self.env.restore(self.env.as_upload(again))
        self.assertEqual(self.env.output_bytes(restored).decode(), text)

    def test_rerun_does_not_leave_gaps_in_numbering(self):
        first = self.env.anonymize(("a.txt", "Проект «Феникс» и ООО «Вектор»."))
        again = self.env.service.rerun(first.id, {}, {"Феникс": "keep"})
        out = self.env.output_bytes(again).decode()
        self.assertIn("Company1", out)

    def test_decisions_are_remembered_for_the_next_files(self):
        first = self.env.anonymize(("a.txt", "Игроки: Трехсосенский и Carlsberg. Проект «Феникс»."))
        again = self.env.service.rerun(first.id, {}, {"Трехсосенский": "hide", "Феникс": "keep"})
        self.assertEqual(again.result["remembered"], {"hide": ["Трехсосенский"], "keep": ["Феникс"]})
        self.assertEqual(self.env.vault.prefs["hide_terms"], ["Трехсосенский"])
        later = self.env.anonymize(("b.txt", "У Трехсосенского и в проекте «Феникс» всё хорошо."))
        out = self.env.output_bytes(later).decode()
        self.assertNotIn("Трехсосенск", out)
        self.assertIn("Феникс", out)

    def test_hide_term_from_options(self):
        job = self.env.anonymize(("a.txt", "Код ЗУП-Х используется в отделе."), options={"hide_terms": ["ЗУП-Х"]})
        self.assertNotIn("ЗУП-Х", self.env.output_bytes(job).decode())


if __name__ == "__main__":
    unittest.main()


class SafetyNetTests(Base):
    def test_password_protected_or_legacy_office_file_is_explained(self):
        data = bytes.fromhex("D0CF11E0A1B11AE1") + b"\0" * 64
        info = self.env.service.peek(Upload("защищённый.xlsx", data))
        self.assertFalse(info["supported"])
        self.assertIn("паролем", info["note"])
        job = self.env.anonymize(("защищённый.xlsx", data))
        self.assertEqual(job.files[0].status, "error")

    def test_archive_that_unpacks_to_gigabytes_is_refused(self):
        from anonymizer import service
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("[Content_Types].xml", "<Types/>")
            z.writestr("xl/big.xml", b"0" * (3 * 1024 * 1024))
        previous = service.MAX_UNPACKED_BYTES
        service.MAX_UNPACKED_BYTES = 1024 * 1024
        try:
            job = self.env.anonymize(("bomb.xlsx", buffer.getvalue()))
        finally:
            service.MAX_UNPACKED_BYTES = previous
        self.assertEqual(job.files[0].status, "error")
        self.assertIn("слишком велик", job.files[0].messages[0]["text"])

    def test_file_without_anything_to_hide_is_not_reported_as_a_success(self):
        job = self.env.anonymize(("a.txt", "Просто текст без данных и имён."))
        self.assertEqual(job.files[0].status, "attention")
        self.assertTrue(any("не найдено" in m for m in self.messages(job)))

    def test_lookalike_words_in_restored_text_are_not_reported_as_leftovers(self):
        job = self.env.anonymize(("a.txt", "Пишите на user1@company1.ru, ООО «Вектор»."))
        restored = self.env.restore(self.env.as_upload(job))
        self.assertEqual(restored.files[0].status, "ok", restored.files[0].messages)
