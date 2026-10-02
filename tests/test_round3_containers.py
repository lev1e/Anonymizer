"""Контейнер Office целиком: эскиз, темы, вложения, служебные части, подсказки, столбцы с географией и перечнями.

Все названия и фамилии здесь выдуманы.
"""
import io
import re
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook

from anonymizer.formats import list_items, place_column_values
from anonymizer.models import Decision, Finding
from anonymizer.service import Service, not_content

from .helpers import Env, content_units, docx_bytes, pptx_bytes, xlsx_bytes, xlsx_cells, zip_text

WEB = Path(__file__).resolve().parent.parent / "anonymizer" / "web" / "index.html"


def rezip(data: bytes, change) -> bytes:
    """Пересобирает контейнер: `change(имя, байты)` возвращает новые байты или None, чтобы убрать часть."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            raw = change(info.filename, zin.read(info.filename))
            if raw is not None:
                zout.writestr(info, raw)
    return out.getvalue()


def add_parts(data: bytes, parts: dict[str, bytes]) -> bytes:
    out = io.BytesIO(rezip(data, lambda n, b: b))
    with zipfile.ZipFile(out, "a") as z:
        for name, raw in parts.items():
            z.writestr(name, raw)
    return out.getvalue()


def names(data_or_path) -> list[str]:
    source = io.BytesIO(data_or_path) if isinstance(data_or_path, bytes) else data_or_path
    with zipfile.ZipFile(source) as z:
        return z.namelist()


def texts(job) -> str:
    return " ".join(m["text"] for f in job.files for m in f.messages)


class Base(unittest.TestCase):
    def setUp(self):
        self.env = Env()

    def tearDown(self):
        self.env.close()

    def round_trip(self, name: str, data: bytes, job):
        """Возврат даёт исходное содержимое по каждой ячейке, абзацу и фигуре."""
        restored = self.env.restore(self.env.as_upload(job))
        self.assertNotEqual(restored.files[0].status, "error", restored.files[0].messages)
        source = Path(self.env.base) / name
        source.write_bytes(data)
        self.assertEqual(content_units(self.env.output(restored)), content_units(source))
        return self.env.output(restored)


class ThumbnailTests(Base):
    """C1: эскиз первой страницы — картинка исходного слайда, его не обезличить, только убрать."""

    def check_dropped(self, name: str, data: bytes):
        self.assertIn("docProps/thumbnail.jpeg", names(data))
        job = self.env.anonymize((name, data))
        out = self.env.output_bytes(job)
        self.assertNotIn("docProps/thumbnail.jpeg", names(out))
        with zipfile.ZipFile(io.BytesIO(out)) as z:
            self.assertNotIn("thumbnail", z.read("_rels/.rels").decode())
            self.assertNotIn("/docProps/thumbnail", z.read("[Content_Types].xml").decode())
        self.assertIn("Эскиз первой страницы удалён", texts(job))
        # Удаление эскиза — не порча структуры: замечаний о целостности нет.
        self.assertNotIn("целостность", texts(job))
        restored = self.round_trip(name, data, job)
        self.assertNotIn("docProps/thumbnail.jpeg", names(restored))

    def test_pptx_thumbnail_is_dropped_with_its_relationship(self):
        self.check_dropped("deck.pptx", pptx_bytes(["Отчёт для ООО «Маяк-Строй» подготовил Иванов Иван Петрович"]))

    def test_docx_thumbnail_is_dropped(self):
        self.check_dropped("memo.docx", docx_bytes(["Письмо ООО «Маяк-Строй» от Иванова Ивана Петровича"]))

    def test_file_without_thumbnail_keeps_every_part(self):
        data = xlsx_bytes({"A1": "Иванов Иван Петрович"})
        job = self.env.anonymize(("book.xlsx", data))
        self.assertEqual(sorted(names(data)), sorted(names(self.env.output_bytes(job))))
        self.assertNotIn("Эскиз", texts(job))

    def test_slide_images_are_kept(self):
        data = add_parts(pptx_bytes(["Слайд"]), {"ppt/media/image9.png": b"\x89PNG fake"})
        job = self.env.anonymize(("pic.pptx", data))
        self.assertIn("ppt/media/image9.png", names(self.env.output_bytes(job)))


class EmbeddedBinaryTests(Base):
    """C1: двоичная книга и объект OLE внутри презентации не разбираются — об этом надо сказать прямо."""

    def test_embedded_xlsb_and_ole_object_raise_a_specific_warning(self):
        data = add_parts(pptx_bytes(["Слайд с таблицей"]), {
            "ppt/embeddings/Microsoft_Excel_Binary_Worksheet.xlsb": b"PK fake binary workbook",
            "ppt/embeddings/oleObject1.bin": b"\xd0\xcf\x11\xe0 fake ole"})
        job = self.env.anonymize(("ole.pptx", data))
        outcome = job.files[0]
        self.assertEqual(outcome.status, "attention")
        warned = [m["text"] for m in outcome.messages if m["level"] == "warn"]
        self.assertTrue(any("встроенный .xlsb" in w and "объект OLE" in w for w in warned), warned)

    def test_embedded_office_package_is_processed_and_not_warned(self):
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        chart = CategoryChartData()
        chart.categories = ["Север", "Юг"]
        chart.add_series("Продажи", (10, 20))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, 4000000, 3000000, chart)
        buffer = io.BytesIO()
        prs.save(buffer)
        job = self.env.anonymize(("chart.pptx", buffer.getvalue()))
        self.assertNotIn("Вложенные объекты не обезличены", texts(job))


def with_theme(data: bytes, theme_name: str, template: str | None = None) -> bytes:
    def change(name, raw):
        if name.startswith("ppt/theme/theme1.xml"):
            text = raw.decode()
            text = text.replace('<a:clrScheme name="Office"', f'<a:clrScheme name="{theme_name} color palette"')
            text = re.sub(r'(<a:theme [^>]*?name=")[^"]*"', rf'\g<1>{theme_name} Theme"', text, count=1)
            return text.encode()
        if name == "docProps/app.xml":
            text = raw.decode().replace("<vt:lpstr>Office Theme</vt:lpstr>", f"<vt:lpstr>{theme_name} Theme</vt:lpstr>")
            if template:
                text = text.replace("<TotalTime>", f"<Template>{template}</Template><TotalTime>")
            return text.encode()
        return raw
    return rezip(data, change)


class ThemeNameTests(Base):
    """C2: имя темы, палитры и шаблона — ярлык автора шаблона, в нём бывает название клиента."""

    def test_theme_palette_template_and_titles_of_parts_are_neutralised_and_restored(self):
        data = with_theme(pptx_bytes(["Итоги квартала"]), "Vantrix", template="Vantrix corporate.potx")
        self.assertIn("Vantrix", zip_text(data))
        job = self.env.anonymize(("theme.pptx", data))
        out = zip_text(self.env.output_bytes(job))
        self.assertNotIn("Vantrix", out)
        # Встроенные имена Office не трогаем: в них нет ничего о владельце.
        self.assertIn('fontScheme name="Office"', out)
        restored = self.round_trip("theme.pptx", data, job)
        back = zip_text(restored.read_bytes())
        self.assertIn('name="Vantrix color palette"', back)
        self.assertIn("<vt:lpstr>Vantrix Theme</vt:lpstr>", back)
        self.assertIn("<Template>Vantrix corporate.potx</Template>", back)

    def test_builtin_names_stay(self):
        data = with_theme(pptx_bytes(["Итоги квартала"]), "Office", template="Normal.dotm")
        job = self.env.anonymize(("plain.pptx", data))
        out = zip_text(self.env.output_bytes(job))
        self.assertIn("<vt:lpstr>Office Theme</vt:lpstr>", out)
        self.assertIn("<Template>Normal.dotm</Template>", out)
        self.assertIn('name="Office Theme"', out)
        self.assertIn('fmtScheme name="Office"', out)


class NonContentSuggestionTests(Base):
    """C3: служебные части файла не дают подсказок «возможно, нужно скрыть»."""

    def test_app_xml_heading_pairs_do_not_become_suggestions(self):
        def change(name, raw):
            if name == "docProps/app.xml":
                return raw.decode().replace("<vt:lpstr>Slide Titles</vt:lpstr>",
                                            "<vt:lpstr>Embedded OLE Servers</vt:lpstr>").encode()
            return raw
        data = rezip(pptx_bytes(["Итоги квартала"]), change)
        job = self.env.anonymize(("app.pptx", data))
        suggested = [s["text"] for s in job.result["suggestions"]]
        self.assertFalse([s for s in suggested if "OLE" in s or "Servers" in s], suggested)

    def test_location_filter(self):
        for location in ("docProps/app.xml::lpstr:3", "ppt/theme/theme2.xml::attribute:clrScheme:name",
                         "ppt/slideMasters/slideMaster1.xml::p:0", "word/styles.xml::attribute:style:name"):
            self.assertTrue(not_content(location), location)
        for location in ("ppt/slides/slide1.xml::p:0", "xl/sharedStrings.xml::si:4", "word/document.xml::p:2",
                         "ppt/slides/slide3.xml::attribute:cNvPr:descr"):
            self.assertFalse(not_content(location), location)


def finding(word: str, location: str, category: str = "POSSIBLE_ENTITY") -> Finding:
    return Finding("x", "f.pptx", location, category, word, 0, len(word), Decision.REVIEW, .5, reason="Похоже на название")


class AltTextTests(unittest.TestCase):
    """C4: в замещающем тексте картинки и подсказке ссылки даже латиница — сильная подсказка."""

    def test_alt_text_candidate_is_strong(self):
        for location in ("ppt/slides/slide1.xml::attribute:cNvPr:descr", "word/document.xml::attribute:docPr:title",
                         "xl/worksheets/sheet1.xml::attribute:hyperlink:tooltip"):
            self.assertTrue(Service._strong_suggestion(finding("Vantrix", location)), location)

    def test_same_word_in_body_text_stays_weak(self):
        self.assertFalse(Service._strong_suggestion(finding("Vantrix", "ppt/slides/slide1.xml::p:0")))
        self.assertFalse(Service._strong_suggestion(finding("ERP", "ppt/slides/slide1.xml::attribute:cNvPr:name")))


class SuggestionCapTests(Base):
    """C5: слабые подсказки не вытесняют сильные, а их общее число видно на экране."""

    def result(self, strong: int, weak: int) -> dict:
        suggestions = {}
        for i in range(strong):
            suggestions[f"s{i}"] = {"text": f"Сильное{i}", "strong": True, "count": 1}
        for i in range(weak):
            suggestions[f"w{i}"] = {"text": f"Weak{i}", "strong": False, "count": 50}
        job = self.env.service.new_job("anonymize")
        from anonymizer.formats import TransformContext
        ctx = TransformContext(self.env.vault, self.env.service.settings_for({}), None)
        return self.env.service._anonymize_result(job, ctx, suggestions)

    def test_many_weak_candidates_do_not_crowd_out_strong_ones(self):
        r = self.result(strong=90, weak=150)
        strong = [s for s in r["suggestions"] if s["strong"]]
        weak = [s for s in r["suggestions"] if not s["strong"]]
        self.assertEqual(len(strong), 90)
        self.assertEqual(len(weak), 60)
        self.assertEqual((r["strong_total"], r["weak_total"]), (90, 150))

    def test_small_lists_are_returned_whole(self):
        r = self.result(strong=2, weak=3)
        self.assertEqual(len(r["suggestions"]), 5)
        self.assertEqual(r["weak_total"], 3)

    def test_banner_names_unreviewed_weak_words(self):
        html = WEB.read_text("utf-8")
        self.assertIn("r.weak_total", html)
        self.assertIn("Не просмотрены менее вероятные слова", html)


def workbook(rows: list[list[str]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class PlaceColumnTests(Base):
    """C6: значения столбцов «Филиал», «Город», «Площадка» — названия мест для всей книги."""

    ROWS = [["Филиал", "Подразделение", "Комментарий"],
            ["Хвойногорье", "Отдел логистики", "План по филиалу Хвойногорье выполнен"],
            ["Заречинский", "Склад", "Заречинский участок работает в две смены"],
            ["Все", "Отдел продаж", "Итого по компании"],
            ["Не указано", "Бухгалтерия", "Нет данных"]]

    def test_place_values_are_hidden_in_cells_and_prose(self):
        data = workbook(self.ROWS)
        job = self.env.anonymize(("places.xlsx", data))
        out = zip_text(self.env.output_bytes(job))
        for place in ("Хвойногорье", "Заречинский"):
            self.assertNotIn(place, out)
        for kept in ("Отдел логистики", "Отдел продаж", "Склад", "Все", "Не указано", "Итого по компании",
                     "План по филиалу", "участок работает"):
            self.assertIn(kept, out)
        self.round_trip("places.xlsx", data, job)

    def test_department_column_is_not_a_place_column(self):
        rows = [["Подразделение", "Отдел"], ["Хвойногорье", "Заречинский"]]
        self.assertEqual(place_column_values({"xl/worksheets/sheet1.xml": _sheet(rows)}), {})

    def test_value_filter(self):
        rows = [["Город", "Регион"], ["г. Верхнеозёрск", "Иркутская область"], ["Иванов И.И.", "Все"],
                ["2024", "головной офис"], ["Санкт-Петербург", "Нет"],
                # Обычные слова из соседней таблицы под тем же столбцом — не места; прилагательное — место.
                ["Выручка", "Дальневосточный"], ["Код", "Выполняется"], ["Частично", "Итого"]]
        values = place_column_values({"xl/worksheets/sheet1.xml": _sheet(rows)})
        self.assertEqual(values, {"г. Верхнеозёрск": "CITY", "Санкт-Петербург": "CITY", "Иркутская область": "REGION",
                                  "Дальневосточный": "REGION"})


class PlaceColumnPrecisionTests(Base):
    """C6: статусы, коды и одиночные значения под «географическим» заголовком — не места."""

    def values(self, rows):
        return place_column_values({"xl/worksheets/sheet1.xml": _sheet(rows)})

    def test_status_column_under_place_like_header_gives_nothing(self):
        rows = [["Регион выполнения", "Площадка"], ["Частично", "Выполняется"], ["Выполняется", "Согласовано"],
                ["Код", "Код"], ["Да", "Проводится"], ["Нет", "Частично"], ["Регион выполнения", "Основной"]]
        self.assertEqual(self.values(rows), {})

    def test_header_text_repeated_below_is_not_a_value(self):
        rows = [["Регион"], ["Регион"], ["Хвойногорье"], ["Заречинский"]]
        self.assertEqual(self.values(rows), {"Хвойногорье": "REGION", "Заречинский": "REGION"})

    def test_single_value_needs_a_geo_cue_elsewhere(self):
        rows = [["Площадка", "Комментарий"], ["Хвойногорье", "Отгрузка по графику"]]
        self.assertEqual(self.values(rows), {})
        rows[1][1] = "Отгрузка в Хвойногорье по графику"
        self.assertEqual(self.values(rows), {"Хвойногорье": "CITY"})

    def test_status_words_stay_in_the_output(self):
        rows = [["Задача", "Регион выполнения", "Площадка"],
                ["Инвентаризация", "Частично", "Выполняется"],
                ["Приёмка", "Выполняется", "Код"],
                ["Отгрузка", "Код", "Частично"]]
        data = workbook(rows)
        job = self.env.anonymize(("status.xlsx", data))
        cells = xlsx_cells(self.env.output_bytes(job))["Sheet"]
        for address, value in (("B2", "Частично"), ("B3", "Выполняется"), ("B4", "Код"), ("C2", "Выполняется"),
                               ("C3", "Код"), ("C4", "Частично")):
            self.assertEqual(cells[address], value, address)


def _sheet(rows):
    from lxml import etree
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    root = etree.Element(f"{{{ns}}}worksheet")
    data = etree.SubElement(root, f"{{{ns}}}sheetData")
    for r, row in enumerate(rows, 1):
        line = etree.SubElement(data, f"{{{ns}}}row")
        for c, value in enumerate(row):
            cell = etree.SubElement(line, f"{{{ns}}}c", r=f"{'ABCDEFG'[c]}{r}", t="inlineStr")
            etree.SubElement(etree.SubElement(cell, f"{{{ns}}}is"), f"{{{ns}}}t").text = value
    return root


class PersonListColumnTests(Base):
    """C6: «А; Б; В» в столбце «Участники» — каждый элемент такой же человек, как одиночное значение столбца."""

    def test_list_items_are_split_after_prefix(self):
        text = "Участники: Сальмаров, Кутеплева; Дерябинцев"
        self.assertEqual([text[a:b] for a, b in list_items(text)], ["Сальмаров", "Кутеплева", "Дерябинцев"])

    def test_every_person_in_a_list_cell_is_hidden_and_restored(self):
        # Полных ФИО в книге нет: узнать фамилии можно только по заголовку столбца. Ячейка с вводным словом
        # («Участники: …») не должна сбивать поиск заголовка.
        rows = [["Совещание", "Участники", "Должность"],
                ["Итоги", "Сальмаров; Кутеплева; Дерябинцев", "Логист, водитель"],
                ["Планёрка", "Тамарчук, Гривцова", "Кладовщик"],
                ["Разбор", "Участники: Шорина, Белькин", "Бухгалтер"]]
        data = workbook(rows)
        job = self.env.anonymize(("people.xlsx", data))
        cells = xlsx_cells(self.env.output_bytes(job))["Sheet"]
        out = " ".join(cells.values())
        for surname in ("Сальмаров", "Кутеплева", "Дерябинцев", "Тамарчук", "Гривцова", "Шорина", "Белькин"):
            self.assertNotIn(surname, out)
        self.assertTrue(cells["B4"].startswith("Участники: "), cells["B4"])
        # Перечень должностей в соседнем столбце — обычные слова.
        self.assertEqual(cells["C2"], "Логист, водитель")
        self.round_trip("people.xlsx", data, job)


class AmountsNoticeTests(Base):
    """C7: числа по умолчанию не заменяются, но о множестве сумм в книге надо сказать."""

    def book(self, amounts: int) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.append(["Статья", "Сумма", "Год"])
        ws.append(["Согласовал Иванов Иван Петрович", 1, 2024])
        for i in range(amounts):
            ws.append([f"Статья {i}", 15320.5 + i * 731, 2024])
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    def test_many_amounts_with_numbers_off_give_a_notice(self):
        job = self.env.anonymize(("money.xlsx", self.book(25)))
        self.assertIn("Заменять числа в Excel", texts(job))
        self.assertEqual(job.files[0].status, "ok")

    def test_no_notice_when_numbers_are_replaced_or_few(self):
        job = self.env.anonymize(("money.xlsx", self.book(25)), options={"numbers": True})
        self.assertNotIn("Числа в таблицах не обезличены", texts(job))
        job = self.env.anonymize(("few.xlsx", self.book(5)))
        self.assertNotIn("Числа в таблицах не обезличены", texts(job))


if __name__ == "__main__":
    unittest.main()
