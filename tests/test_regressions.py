"""Дефекты, найденные при ручной проверке расширенных файлов и в окне программы. Каждый закреплён проверкой."""
import io
import re
import unittest
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from .helpers import Env


def anonymize_text(text: str, env: Env | None = None) -> str:
    env = env or Env()
    job = env.anonymize(("t.txt", text))
    return Path(job.files[0].out_path).read_text("utf-8")


class EmailAndPhoneTests(unittest.TestCase):
    def test_email_before_a_full_stop_is_hidden_whole(self):
        """Точка в конце предложения не входит в адрес; раньше адрес с ней не находился, а домен заменялся отдельно."""
        for text in ("Пишите: kuznetsova@technopark-n.ru.", "(kuznetsova@technopark-n.ru).", "Почта kuznetsova@technopark-n.ru, тел.",
                     "Почта: kuznetsova@technopark-n.ru...", "e-mail kuznetsova@technopark-n.ru!\nДалее"):
            with self.subTest(text=text):
                out = anonymize_text(text)
                self.assertNotIn("kuznetsova", out)
                self.assertNotIn("technopark", out)
                self.assertRegex(out, r"email\d+@example\.com")

    def test_email_after_a_company_with_a_similar_name(self):
        out = anonymize_text("АО «Технопарк». Контакт: Кузнецова М.С., e-mail: kuznetsova@technopark-n.ru.")
        self.assertNotIn("kuznetsova", out)
        self.assertNotIn("Технопарк", out)

    def test_dotted_local_part_and_subdomain(self):
        out = anonymize_text("Пишите a.smirnov@mail.mayak-stroy.ru. Спасибо.")
        self.assertNotIn("smirnov", out)
        self.assertNotIn("mayak", out)

    def test_international_phones(self):
        for phone in ("+44 20 7946 0958", "+1 (415) 555-2671", "+49 30 1234567", "+7 495 123-45-67", "+33 1 42 68 53 00",
                      "+380 44 123 4567", "+86 10 6552 9988"):
            with self.subTest(phone=phone):
                out = anonymize_text(f"Call {phone} today")
                self.assertNotIn(phone.split()[-1], out)
                self.assertRegex(out, r"Phone\d+")

    def test_numbers_that_look_like_growth_or_money_are_not_phones(self):
        for text in ("Рост составил +12 345 678 руб.", "Изменение +5% к прошлому году", "Отклонение +1 250 000 000 ₽", "Итого 2026-03-15"):
            with self.subTest(text=text):
                self.assertNotIn("Phone", anonymize_text(text))


class AddressTests(unittest.TestCase):
    def test_street_with_a_house_number_but_no_marker(self):
        out = anonymize_text("Приезжайте в офис на ул. Мира, 7, Красноярск.")
        self.assertNotIn("Мира", out)

    def test_street_name_alone(self):
        for text, word in (("Офис на ул. Ленина.", "Ленина"), ("Склад: улица Большая Садовая", "Садовая"),
                           ("Адрес доставки — пер. Гранитный", "Гранитный")):
            with self.subTest(text=text):
                self.assertNotIn(word, anonymize_text(text))

    def test_ordinary_words_after_street_abbreviations_survive(self):
        out = anonymize_text("Скорость на улице выросла. Улица длиннее переулка. Пер. 2 кв. вырос.")
        self.assertNotIn("Address", out)


class PersonTests(unittest.TestCase):
    def test_honorific_before_a_surname(self):
        out = anonymize_text("Dear Mr. Johnson, please call Dr. Watson and Mrs. Hudson.")
        for word in ("Johnson", "Watson", "Hudson"):
            self.assertNotIn(word, out)
        self.assertIn("Mr.", out)

    def test_same_honorific_surname_gets_the_same_label(self):
        out = anonymize_text("Mr. Johnson arrived. Later Mr. Johnson left. Then Johnson called.")
        self.assertNotIn("Johnson", out)
        self.assertEqual(len(set(re.findall(r"Name\d+", out))), 1, out)

    def test_russian_honorifics(self):
        out = anonymize_text("Передайте г-ну Ковалёву, г-же Лаптевой и господину Сидорову.")
        for word in ("Ковалёв", "Лаптев", "Сидоров"):
            self.assertNotIn(word, out)

    def test_given_and_patronymic_is_one_person_and_not_glued_to_someone_else(self):
        """Ольга Николаевна из письма — не Морозова Ольга Викторовна из другого файла, и отчество не остаётся в файле."""
        env = Env()
        job = env.anonymize(("a.txt", "Бюджет согласован с финансовым директором, Морозовой Ольгой Викторовной."),
                            ("b.txt", "Добрый день, Ольга Николаевна! Звоните Ольге Николаевне."))
        out = Path(job.files[1].out_path).read_text("utf-8")
        self.assertNotIn("Николаев", out)
        self.assertNotIn("Ольг", out)
        first = Path(job.files[0].out_path).read_text("utf-8")
        self.assertNotEqual(re.findall(r"Name\d+", first)[0], re.findall(r"Name\d+", out)[0])

    def test_name_at_the_end_of_a_line_before_a_capitalised_line(self):
        """Перевод строки разделяет записи: «John Smith» и следующая строка «Senior Analyst» — не название из трёх слов."""
        out = anonymize_text("Best regards,\nJohn Smith\nSenior Analyst, Northwind Traders\n")
        self.assertNotIn("Smith", out)
        self.assertNotIn("John", out)


class OrganisationTests(unittest.TestCase):
    def test_region_word_inside_a_quoted_company_name_does_not_split_it(self):
        for text, rest in (("Акционерное общество «Технопарк Север», именуемое Заказчик", "Технопарк"),
                           ("ООО «Приволжский Банк Развития» подписало", "Банк Развития"),
                           ("ПАО «Газпром Нефть Москва» сообщило", "Газпром")):
            with self.subTest(text=text):
                out = anonymize_text(text)
                self.assertNotIn(rest, out)
                self.assertRegex(out, r"«Company\d+»")

    def test_object_in_the_dative_is_a_project(self):
        out = anonymize_text("Работы по объекту «Порог-2» завершены, по проекту «Атлант» продолжаются.")
        self.assertEqual(len(re.findall(r"Project\d+", out)), 2, out)
        self.assertNotIn("Company", out)


class MetadataTests(unittest.TestCase):
    """Название, тема, ключевые слова и категория документа — те же данные, что и текст: в них пишут имя клиента."""

    def _make(self, kind: str) -> bytes:
        buffer = io.BytesIO()
        if kind == "docx":
            doc = Document()
            doc.add_paragraph("Текст")
            props = doc.core_properties
        elif kind == "xlsx":
            wb = Workbook()
            wb.active["A1"] = "текст"
            props = wb.properties
        else:
            doc = Presentation()
            doc.slides.add_slide(doc.slide_layouts[6])
            props = doc.core_properties
        props.title = "Отчёт для ООО «Маяк-Строй»"
        props.subject = "Проект Порог"
        props.keywords = "Смирнов, Маяк-Строй"
        props.category = "Клиент Технопарк"
        if kind == "xlsx":
            props.creator, props.description = "Смирнов Алексей Петрович", "Кузнецова Мария Сергеевна"
            wb.save(buffer)
        else:
            props.author, props.comments = "Смирнов Алексей Петрович", "Кузнецова Мария Сергеевна"
            doc.save(buffer)
        return buffer.getvalue()

    def test_document_properties_are_hidden_and_come_back(self):
        for kind in ("docx", "xlsx", "pptx"):
            with self.subTest(kind=kind):
                env = Env()
                job = env.anonymize((f"a.{kind}", self._make(kind)))
                core = zipfile.ZipFile(job.files[0].out_path).read("docProps/core.xml").decode()
                for word in ("Маяк", "Смирнов", "Кузнецов", "Порог", "Технопарк"):
                    self.assertNotIn(word, core)
                back = env.restore(env.as_upload(job))
                restored = zipfile.ZipFile(back.files[0].out_path).read("docProps/core.xml").decode()
                self.assertIn("Отчёт для ООО «Маяк-Строй»", restored)
                self.assertIn("Клиент Технопарк", restored)

    def test_company_and_manager_in_app_properties(self):
        env = Env()
        doc = Document()
        doc.add_paragraph("Текст")
        buffer = io.BytesIO()
        doc.save(buffer)
        source = zipfile.ZipFile(io.BytesIO(buffer.getvalue()))
        patched = io.BytesIO()
        with zipfile.ZipFile(patched, "w", zipfile.ZIP_DEFLATED) as z:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "docProps/app.xml":
                    data = data.decode("utf-8").replace(
                        "</Properties>", "<Company>ООО «Маяк-Строй»</Company><Manager>Смирнов Алексей Петрович</Manager></Properties>").encode()
                z.writestr(item, data)
        job = env.anonymize(("a.docx", patched.getvalue()))
        app = zipfile.ZipFile(job.files[0].out_path).read("docProps/app.xml").decode()
        self.assertNotIn("Маяк", app)
        self.assertNotIn("Смирнов", app)



class ManualReviewOfCaseFiles(unittest.TestCase):
    """Найдено при ручном просмотре списка замен в файлах case: служебные строки принимались за названия компаний."""

    def test_version_encoding_and_plain_numbers_in_quotes_are_not_companies(self):
        out = anonymize_text('Версия "1.0", кодировка "UTF-16", код "32687", шаблон "0.00", файл «38530».')
        self.assertNotIn("Company", out)

    def test_powerpoint_addin_payload_in_a_tag_is_left_alone(self):
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        buffer = io.BytesIO()
        prs.save(buffer)
        src = zipfile.ZipFile(io.BytesIO(buffer.getvalue()))
        out = io.BytesIO()
        payload = "&lt;?xml version=&quot;1.0&quot; encoding=&quot;UTF-16&quot;?&gt;&lt;root name=&quot;ООО «Ромашка»&quot;/&gt;"
        with zipfile.ZipFile(out, "w") as z:
            for item in src.infolist():
                z.writestr(item, src.read(item.filename))
            z.writestr("ppt/tags/tag1.xml", '<p:tagLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
                                            f'<p:tag name="ADDIN" val="{payload}"/></p:tagLst>')
        env = Env()
        job = env.anonymize(("a.pptx", out.getvalue()))
        tag = raw_parts(job.files[0].out_path, "ppt/tags/tag1.xml")
        self.assertIn("UTF-16", tag)
        self.assertIn("Ромашка", tag)
        self.assertNotIn("Company", tag)


def raw_parts(path, *names) -> str:
    with zipfile.ZipFile(path) as z:
        return "\n".join(z.read(n).decode("utf-8", "ignore") for n in names if n in z.namelist())


if __name__ == "__main__":
    unittest.main()
