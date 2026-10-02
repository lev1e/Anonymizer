"""Общие заготовки для проверок: изолированное хранилище и сборка небольших файлов."""
from __future__ import annotations

import io
import os
import tempfile
import zipfile
from pathlib import Path

import fitz
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches

from anonymizer.service import Service, Upload
from anonymizer.vault import Vault

CASE_DIR = Path(__file__).resolve().parent.parent / "case"


class Env:
    """Хранилище и сервис во временной папке: проверки не трогают настоящие данные пользователя."""

    def __init__(self, **prefs):
        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)   # Windows держит открытые PDF до сборки мусора
        self.base = Path(self.tmp.name)
        self.key = os.urandom(32)
        self.vault = Vault(self.base / "vault.bin", key=self.key)
        for name, value in prefs.items():
            self.vault.prefs[name] = value
        self.service = Service(self.vault, self.base / "work")

    def close(self) -> None:
        self.tmp.cleanup()

    def reopen(self) -> "Env":
        """Тот же диск, новый процесс: проверяет, что всё нужное сохранилось."""
        self.vault = Vault(self.base / "vault.bin", key=self.key)
        self.service = Service(self.vault, self.base / "work")
        return self

    @staticmethod
    def upload(item) -> Upload:
        if isinstance(item, Upload):
            return item
        if isinstance(item, Path):
            return Upload(item.name, item.read_bytes())
        name, data = item
        return Upload(name, data if isinstance(data, bytes) else data.encode("utf-8"))

    def anonymize(self, *items, options=None, overrides=None):
        return self.service.anonymize([self.upload(i) for i in items], options or {}, overrides)

    def restore(self, *items):
        return self.service.restore([self.upload(i) for i in items])

    @staticmethod
    def output(job, index: int = 0) -> Path:
        return Path(job.files[index].out_path)

    @staticmethod
    def output_bytes(job, index: int = 0) -> bytes:
        return Path(job.files[index].out_path).read_bytes()

    @staticmethod
    def as_upload(job, index: int = 0, name: str | None = None) -> Upload:
        path = Path(job.files[index].out_path)
        return Upload(name or path.name, path.read_bytes())


def xlsx_bytes(cells: dict[str, object], title: str = "Лист1", extra_sheets: dict[str, dict] | None = None,
               creator: str | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for address, value in cells.items():
        ws[address] = value
    for name, sheet_cells in (extra_sheets or {}).items():
        sheet = wb.create_sheet(name)
        for address, value in sheet_cells.items():
            sheet[address] = value
    if creator:
        wb.properties.creator = creator
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def docx_bytes(paragraphs: list[str], table: list[list[str]] | None = None, author: str | None = None) -> bytes:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table:
        grid = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    if author:
        doc.core_properties.author = author
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def pptx_bytes(texts: list[str], title: str | None = None) -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    if title is not None:
        slide.shapes.title.text = title
    for index, text in enumerate(texts):
        box = slide.shapes.add_textbox(Inches(1), Inches(1.5 + index), Inches(6), Inches(.8))
        box.text_frame.text = text
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def pdf_bytes(text: str, cyrillic: bool = False) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text, fontname="china-s" if cyrillic else "helv", fontsize=12)
    data = doc.tobytes()
    doc.close()
    return data


def xlsx_cells(path_or_bytes) -> dict[str, dict[str, object]]:
    source = io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, bytes) else path_or_bytes
    wb = load_workbook(source)
    return {ws.title: {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
            for ws in wb.worksheets}


def docx_text(path_or_bytes) -> str:
    source = io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, bytes) else str(path_or_bytes)
    doc = Document(source)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def zip_text(path_or_bytes) -> str:
    """Весь XML контейнера одной строкой — для проверки, что имя нигде не осталось."""
    source = io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, bytes) else path_or_bytes
    with zipfile.ZipFile(source) as z:
        return "\n".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist()
                         if n.endswith((".xml", ".rels")))


def content_units(path: Path) -> dict[str, str]:
    """Содержимое файла по адресам: ячейка, абзац, фигура. Так сравниваются исходник и восстановленный файл."""
    from openpyxl import load_workbook as _load
    units: dict[str, str] = {}
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        wb = _load(path)
        for ws in wb.worksheets:
            units[f"sheet:{ws.title}"] = ws.title
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        units[f"{ws.title}!{cell.coordinate}"] = str(cell.value)
    elif suffix == ".docx":
        doc = Document(str(path))
        for index, paragraph in enumerate(doc.paragraphs):
            units[f"p{index}"] = paragraph.text
        for t_index, table in enumerate(doc.tables):
            for r_index, row in enumerate(table.rows):
                for c_index, cell in enumerate(row.cells):
                    units[f"t{t_index}.{r_index}.{c_index}"] = cell.text
    elif suffix == ".pptx":
        prs = Presentation(str(path))

        def walk(shapes, prefix, slide):
            for shape in shapes:
                if shape.shape_type == 6:
                    walk(shape.shapes, prefix + shape.name + "/", slide)
                if shape.has_text_frame:
                    units[f"s{slide}:{prefix}{shape.name}"] = shape.text_frame.text
                if getattr(shape, "has_table", False) and shape.has_table:
                    for r_index, row in enumerate(shape.table.rows):
                        for c_index, cell in enumerate(row.cells):
                            units[f"s{slide}:{shape.name}.{r_index}.{c_index}"] = cell.text
        for number, slide in enumerate(prs.slides, 1):
            walk(slide.shapes, "", number)
    return units
