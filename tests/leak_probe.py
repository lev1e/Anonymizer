"""Сложные документы с секретами в каждом укромном месте. Общие помощники для проверок на утечки.

Секрет считается утёкшим, если он читается в любой части файла: в тексте, в атрибуте, в имени части, в рисунке-подсказке."""
import io
import re
import zipfile

from docx import Document
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

SECRETS = ["Смирнов", "Алексей", "Ромашк", "913 555", "romashka", "Новосибирск", "Иванова"]
NAME = "Смирнов Алексей Петрович"
CLIENT = "ООО «Ромашка»"


def everything(data: bytes) -> dict[str, str]:
    """Весь читаемый материал файла по частям: текст, значения атрибутов, имя части."""
    out = {}
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for name in z.namelist():
            raw = z.read(name)
            if name.endswith((".xml", ".rels", ".vml")):
                text = raw.decode("utf-8", "ignore")
                attrs = " ".join(re.findall(r'="([^"]*)"', text))
                out[name] = name + "\n" + re.sub(r"<[^>]+>", " ", text) + "\n" + attrs
            else:
                out[name] = name
    return out


def leaks(data: bytes, secrets=SECRETS) -> list[tuple[str, str]]:
    found = []
    for part, text in everything(data).items():
        low = text.lower()
        for secret in secrets:
            if secret.lower() in low:
                i = low.find(secret.lower())
                found.append((part, text[max(0, i - 40):i + 50].replace("\n", " ")))
    return found


def inject(data: bytes, parts: dict[str, bytes | None], patches: dict | None = None) -> bytes:
    """Добавляет или заменяет части файла; patches = {часть: функция(текст) → текст}."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        seen = set()
        for info in src.infolist():
            blob = src.read(info.filename)
            if info.filename in parts:
                if parts[info.filename] is None:
                    continue
                blob = parts[info.filename]
            if patches and info.filename in patches:
                blob = patches[info.filename](blob.decode("utf-8")).encode("utf-8")
            dst.writestr(info, blob)
            seen.add(info.filename)
        for name, blob in parts.items():
            if name not in seen and blob is not None:
                dst.writestr(name, blob)
    return out.getvalue()


def add_override(content_types: str, part: str, ctype: str) -> str:
    return content_types.replace("</Types>", f'<Override PartName="/{part}" ContentType="{ctype}"/></Types>')


# -- Word ----------------------------------------------------------------------------

def rich_docx() -> bytes:
    d = Document()
    d.core_properties.author = NAME
    d.core_properties.title = f"Договор {CLIENT}"
    d.core_properties.subject = "Поставка для Ромашка"
    d.core_properties.keywords = "Ромашка; Смирнов"
    d.core_properties.comments = f"Составил {NAME}"
    d.sections[0].header.paragraphs[0].text = f"{CLIENT} — конфиденциально"
    d.sections[0].footer.paragraphs[0].text = f"{NAME}, +7 913 555-12-34"
    d.sections[0].different_first_page_header_footer = True
    d.sections[0].first_page_header.paragraphs[0].text = f"Титул {CLIENT}"
    d.add_heading(f"Договор поставки {CLIENT}", 1)
    p = d.add_paragraph(f"Директор {NAME}, г. Новосибирск, сайт romashka.ru.")
    p.add_run(" Позвоните помощнику Ивановой М.С.")
    p2 = d.add_paragraph("Ссылка: ")
    # внешняя ссылка с доменом клиента
    part = d.part
    rid = part.relate_to("https://portal.romashka.ru/смирнов", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    from docx.oxml import parse_xml
    h = parse_xml(f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                  f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="{rid}">'
                  f'<w:r><w:t>портал Ромашка</w:t></w:r></w:hyperlink>')
    p2._p.append(h)
    d.add_comment(p.runs[0], text=f"Проверить у {NAME}", author="Смирнов А.П.", initials="СА")
    out = io.BytesIO()
    d.save(out)
    data = out.getvalue()

    footnotes = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                 '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
                 '<w:footnote w:id="1"><w:p><w:r><w:t>Сноска: подпись Смирнов Алексей Петрович, ООО «Ромашка».</w:t></w:r></w:p></w:footnote></w:footnotes>')
    sdt = ('<w:sdt><w:sdtPr><w:alias w:val="Клиент Ромашка"/><w:tag w:val="romashka_contact"/></w:sdtPr>'
           '<w:sdtContent><w:p><w:r><w:t>Контакт: Смирнов Алексей Петрович</w:t></w:r></w:p></w:sdtContent></w:sdt>')
    textbox = ('<w:p><w:r><w:pict><v:shape xmlns:v="urn:schemas-microsoft-com:vml"><v:textbox><w:txbxContent>'
               '<w:p><w:r><w:t>В рамке: ООО «Ромашка», Смирнов</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>')
    bookmark = '<w:p><w:bookmarkStart w:id="9" w:name="Ромашка_раздел"/><w:r><w:t>Раздел</w:t></w:r><w:bookmarkEnd w:id="9"/></w:p>'
    field = ('<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> HYPERLINK "https://romashka.ru/смирнов" </w:instrText></w:r>'
             '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>ссылка</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>')
    custom = ('<?xml version="1.0" encoding="UTF-8"?><Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
              'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"><property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="Клиент">'
              '<vt:lpwstr>ООО «Ромашка»</vt:lpwstr></property></Properties>')

    def patch_doc(xml: str) -> str:
        return xml.replace("<w:sectPr", sdt + textbox + bookmark + field + "<w:sectPr", 1)

    def patch_ct(xml: str) -> str:
        xml = add_override(xml, "word/footnotes.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml")
        return add_override(xml, "docProps/custom.xml", "application/vnd.openxmlformats-officedocument.custom-properties+xml")

    def patch_rels(xml: str) -> str:
        return xml.replace("</Relationships>", '<Relationship Id="rIdFn" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>')

    def patch_root_rels(xml: str) -> str:
        return xml.replace("</Relationships>", '<Relationship Id="rIdCp" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties" Target="docProps/custom.xml"/></Relationships>')

    return inject(data, {"word/footnotes.xml": footnotes.encode(), "docProps/custom.xml": custom.encode()},
                  {"word/document.xml": patch_doc, "[Content_Types].xml": patch_ct, "word/_rels/document.xml.rels": patch_rels,
                   "_rels/.rels": patch_root_rels})


# -- Excel ---------------------------------------------------------------------------

def rich_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Реестр"
    ws.append(["ФИО", "Компания", "Город", "Сумма"])
    ws.append([NAME, CLIENT, "Новосибирск", 1200])
    ws.append(["Иванова Мария Сергеевна", CLIENT, "Новосибирск", 800])
    ws["A2"].comment = Comment(f"Звонил {NAME}", "Смирнов А.П.")
    ws.oddHeader.center.text = CLIENT
    ws.oddFooter.left.text = NAME
    dv = DataValidation(type="list", formula1='"Ромашка,Лютик"', showErrorMessage=True, error="Только Ромашка", errorTitle="Клиент")
    ws.add_data_validation(dv)
    dv.add("B2:B3")
    hidden = wb.create_sheet("Скрытый Ромашка")
    hidden["A1"] = f"Секрет: {NAME}, +7 913 555-12-34"
    hidden.sheet_state = "hidden"
    very = wb.create_sheet("Служебный")
    very["A1"] = "romashka.ru"
    very.sheet_state = "veryHidden"
    chart = BarChart()
    chart.title = f"Выручка {CLIENT}"
    chart.add_data(Reference(ws, min_col=4, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=3))
    ws.add_chart(chart, "F2")
    wb.defined_names["Клиент_Ромашка"] = DefinedName("Клиент_Ромашка", attr_text="Реестр!$B$2")
    wb.properties.creator = NAME
    wb.properties.title = f"Реестр {CLIENT}"
    wb.properties.keywords = "Ромашка"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# -- PowerPoint ----------------------------------------------------------------------

def rich_pptx() -> bytes:
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = f"Проект для {CLIENT}"
    s.placeholders[1].text = f"Руководитель: {NAME}\nГород: Новосибирск"
    s.notes_slide.notes_text_frame.text = f"Заметка: позвонить {NAME}, +7 913 555-12-34"
    pic_shape = s.shapes.add_textbox(Inches(1), Inches(5), Inches(3), Inches(1))
    pic_shape.text_frame.text = "Подпись Иванова М.С."
    pic_shape.name = "Блок Ромашка"
    chart_data = CategoryChartData()
    chart_data.categories = ["Смирнов", "Иванова"]
    chart_data.add_series("Выручка Ромашка", (1.2, 0.8))
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Итоги"
    s2.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(6), Inches(4), chart_data)
    prs.core_properties.author = NAME
    prs.core_properties.title = f"Презентация {CLIENT}"
    out = io.BytesIO()
    prs.save(out)

    def shape(text: str) -> str:
        return ('<p:sp><p:nvSpPr><p:cNvPr id="99" name="Логотип Ромашка" descr="Логотип ООО Ромашка"/><p:cNvSpPr txBox="1"/><p:nvPr userDrawn="1"/></p:nvSpPr>'
                '<p:spPr><a:xfrm><a:off x="457200" y="6400800"/><a:ext cx="4572000" cy="365760"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>'
                f'<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="ru-RU"/><a:t>{text}</a:t></a:r></a:p></p:txBody></p:sp>')

    # образец и макет слайдов с названием клиента: так выглядят корпоративные шаблоны
    def into_tree(text: str):
        return lambda xml: xml.replace("</p:spTree>", shape(text) + "</p:spTree>", 1)
    return inject(out.getvalue(), {}, {"ppt/slideMasters/slideMaster1.xml": into_tree("ООО «Ромашка» — конфиденциально"),
                                       "ppt/slideLayouts/slideLayout2.xml": into_tree("Шаблон Ромашка, Смирнов А.П.")})
