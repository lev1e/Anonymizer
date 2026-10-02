"""Add-in: сертификаты, защита сервера, полный круг «обезличить → проверить → вернуть → откатить» по HTTPS."""
import http.client
import json
import ssl
import threading
import time
import unittest
import urllib.parse
from pathlib import Path

from anonymizer import addin, certs
from anonymizer.addin import Originals, create_addin_server, structure_of, verify_document
from anonymizer.webapp import App

from .helpers import Env, docx_bytes, docx_text, pptx_bytes, xlsx_bytes, xlsx_cells

SECRET = "Смирнов Алексей Петрович"


class AddinCase(unittest.TestCase):
    def setUp(self):
        self.env = Env()
        self.app = App(self.env.vault, self.env.base / "work", "tok")
        self.dir = self.env.base / "addin"
        self.files = certs.ensure_certs(self.dir, self.env.key)
        self.server, self.api = create_addin_server(self.app, port=0, files=self.files, directory=self.dir, key=self.env.key)
        self.port = self.server.server_address[1]
        threading.Thread(target=self.server.serve_forever, daemon=True).start()
        self.ctx = ssl.create_default_context(cafile=str(self.files.ca_cert))

    def tearDown(self):
        self.server.shutdown()
        self.server.server_close()
        self.env.close()

    # -- запросы ---------------------------------------------------------------

    def call(self, method, path, body=None, headers=None, token=True, host=None, auth=True):
        conn = http.client.HTTPSConnection("localhost", self.port, context=self.ctx, timeout=60)
        head = {"Host": host or f"localhost:{self.port}"}
        if auth:
            head.update({"X-Requested-With": "anonymizer"})
            if token:
                head["X-Addin-Token"] = self.api.token
        head.update(headers or {})
        conn.request(method, path, body=body, headers=head)
        response = conn.getresponse()
        data = response.read()
        conn.close()
        return response.status, dict(response.getheaders()), data

    def json(self, *args, **kwargs):
        status, _, data = self.call(*args, **kwargs)
        return status, json.loads(data.decode("utf-8")) if data else {}

    def run_op(self, kind, name, data, options=None):
        status, started = self.json("POST", f"/api/addin/{kind}", data, {
            "X-Filename": urllib.parse.quote(name), "X-Options": urllib.parse.quote(json.dumps(options or {}))})
        self.assertEqual(status, 200, started)
        for _ in range(600):
            status, job = self.json("GET", f"/api/addin/jobs/{started['job']}")
            if job["state"] != "running":
                return started, job
            time.sleep(0.1)
        self.fail("задание не завершилось")

    def result(self, job_id):
        status, _, data = self.call("GET", f"/api/addin/result/{job_id}")
        self.assertEqual(status, 200)
        return data


class CertTests(AddinCase):
    def test_server_certificate_is_valid_for_localhost_and_ip(self):
        status, _, _ = self.call("GET", "/addin/taskpane.html", auth=False)
        self.assertEqual(status, 200)
        context = ssl.create_default_context(cafile=str(self.files.ca_cert))
        conn = http.client.HTTPSConnection("127.0.0.1", self.port, context=context, timeout=10)
        conn.request("GET", "/addin/taskpane.html", headers={"Host": f"127.0.0.1:{self.port}"})
        self.assertEqual(conn.getresponse().status, 200)

    def test_ca_is_limited_to_localhost(self):
        from cryptography import x509
        ca = x509.load_pem_x509_certificate(self.files.ca_cert.read_bytes())
        constraints = ca.extensions.get_extension_for_class(x509.NameConstraints)
        self.assertTrue(constraints.critical)
        self.assertEqual([n.value for n in constraints.value.permitted_subtrees if isinstance(n, x509.DNSName)], ["localhost"])

    def test_certificates_are_reused_and_private_keys_encrypted(self):
        again = certs.ensure_certs(self.dir, self.env.key)
        self.assertFalse(again.ca_created or again.leaf_created)
        self.assertIn(b"ENCRYPTED", self.files.leaf_key.read_bytes())
        self.assertIn(b"ENCRYPTED", self.files.ca_key.read_bytes())

    def test_expiring_leaf_is_renewed_without_new_ca(self):
        import datetime as dt
        old_ca = self.files.ca_cert.read_bytes()
        future = dt.datetime.now(dt.timezone.utc) + dt.timedelta(days=certs.LEAF_DAYS - 10)
        renewed = certs.ensure_certs(self.dir, self.env.key, now=future)
        self.assertTrue(renewed.leaf_created)
        self.assertFalse(renewed.ca_created)
        self.assertEqual(self.files.ca_cert.read_bytes(), old_ca)


class SecurityTests(AddinCase):
    def test_api_requires_token_and_header(self):
        self.assertEqual(self.call("GET", "/api/addin/info", token=False)[0], 403)
        self.assertEqual(self.call("GET", "/api/addin/info", auth=False)[0], 403)
        self.assertEqual(self.call("GET", "/api/addin/info")[0], 200)

    def test_foreign_origin_and_host_are_rejected(self):
        self.assertEqual(self.call("GET", "/api/addin/info", headers={"Origin": "https://evil.example"})[0], 403)
        self.assertEqual(self.call("GET", "/api/addin/info", headers={"Sec-Fetch-Site": "cross-site"})[0], 403)
        self.assertEqual(self.call("GET", "/api/addin/info", host="evil.example")[0], 403)
        self.assertEqual(self.call("GET", "/api/addin/info", headers={"Origin": f"https://localhost:{self.port}"})[0], 200)

    def test_panel_server_does_not_expose_window_only_routes(self):
        for path in ("/api/native/upload-path", "/api/native/folder", "/api/quit", "/api/backup/export", "/api/vault/clear", "/api/prefs"):
            status, _, _ = self.call("POST", path, b"{}")
            self.assertEqual(status, 404, path)
        self.assertTrue(self.app.windows == {} and not self.app.quit_event.is_set())

    def test_post_without_token_does_not_start_work(self):
        status, _, _ = self.call("POST", "/api/addin/anonymize", b"x", {"X-Filename": "a.docx"}, token=False)
        self.assertEqual(status, 403)
        self.assertEqual(self.app.uploads, {})

    def test_static_does_not_escape_directory_or_serve_templates(self):
        for path in ("/addin/../addin.py", "/addin/%2e%2e/addin.py", "/addin/manifest.xml.tmpl", "/addin/vendor/../../vault.py"):
            self.assertEqual(self.call("GET", path, auth=False)[0], 404, path)

    def test_page_carries_the_token_and_vendor_runtime_is_local(self):
        status, _, data = self.call("GET", "/addin/taskpane.html", auth=False)
        text = data.decode("utf-8")
        self.assertIn(self.api.token, text)
        self.assertNotIn("__ADDIN_TOKEN__", text)
        self.assertNotIn("appsforoffice.microsoft.com", text)
        from anonymizer.addin import ADDIN_DIR
        if (ADDIN_DIR / "vendor" / "office.js").is_file():          # скачивается tools/fetch_office_js.py
            self.assertEqual(self.call("GET", "/addin/vendor/office.js", auth=False)[0], 200)


class FlowTests(AddinCase):
    def test_docx_roundtrip_with_original_copy_and_verification(self):
        import io

        from docx import Document
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = "ООО «Ромашка» — конфиденциально"
        doc.add_paragraph(f"Директор {SECRET}, ООО «Ромашка».")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "ФИО", SECRET
        buffer = io.BytesIO()
        doc.save(buffer)
        original = buffer.getvalue()
        started, job = self.run_op("anonymize", "Договор.docx", original)
        self.assertEqual(job["state"], "done", job)
        self.assertTrue(job["files"][0]["downloadable"])
        hidden = self.result(job["id"])
        self.assertNotIn("Смирнов", docx_text(hidden))
        # Копия оригинала хранится зашифрованной и возвращается ровно такой же.
        status, headers, copy = self.call("GET", f"/api/addin/backup/{started['backup']}")
        self.assertEqual((status, copy), (200, original))
        status, info = self.json("GET", f"/api/addin/backup-info/{started['backup']}")
        self.assertEqual(status, 200)
        self.assertTrue(any(p["kind"] == "header" for p in info["structure"]["parts"]))     # колонтитулы приходят телом ответа
        self.assertNotIn(original[:200], b"".join(p.read_bytes() for p in (self.dir / "originals").iterdir()))
        # Проверка после замены: тот самый файл проходит, исходный (с данными) нет.
        status, verdict = self.json("POST", f"/api/addin/verify/{job['id']}", hidden)
        self.assertTrue(verdict["ok"], verdict)
        status, verdict = self.json("POST", f"/api/addin/verify/{job['id']}", original)
        self.assertFalse(verdict["ok"])
        self.assertIn("leak", {p["code"] for p in verdict["problems"]})
        # Возврат
        _, back = self.run_op("restore", "Договор.docx", hidden)
        self.assertEqual(back["state"], "done", back)
        self.assertIn(SECRET, docx_text(self.result(back["id"])))

    def test_edits_made_during_processing_are_detected_before_replacing(self):
        original = docx_bytes([f"Директор {SECRET}.", "Абзац про поставки."])
        _, job = self.run_op("anonymize", "a.docx", original)
        status, same = self.json("POST", f"/api/addin/unchanged/{job['id']}", original)
        self.assertEqual((status, same["same"]), (200, True), same)
        edited = docx_bytes([f"Директор {SECRET}.", "Абзац про поставки.", "Человек дописал новый абзац."])
        status, changed = self.json("POST", f"/api/addin/unchanged/{job['id']}", edited)
        self.assertFalse(changed["same"])
        self.assertGreater(changed["changed"], 0)

    def test_lost_text_is_reported(self):
        _, job = self.run_op("anonymize", "a.docx", docx_bytes([f"Директор {SECRET}.", "Второй абзац про поставки."]))
        shorter = docx_bytes([docx_text(self.result(job["id"])).splitlines()[0]])
        status, verdict = self.json("POST", f"/api/addin/verify/{job['id']}", shorter)
        self.assertIn("lost_text", {p["code"] for p in verdict["problems"]})

    def test_xlsx_structure_is_returned_for_rebuilding_the_workbook(self):
        data = xlsx_bytes({"A1": "ФИО", "A2": SECRET}, title=SECRET, extra_sheets={"Итоги": {"A1": 1}})
        _, job = self.run_op("anonymize", "Реестр.xlsx", data)
        self.assertEqual(job["state"], "done", job)
        names = [s["name"] for s in job["structure"]["sheets"]]
        self.assertEqual(len(names), 2)
        self.assertNotIn("Смирнов", " ".join(names))
        self.assertIn("Итоги", names)

    def test_pptx_slide_count_is_checked(self):
        _, job = self.run_op("anonymize", "Показ.pptx", pptx_bytes([f"Директор {SECRET}", "Второй слайд"]))
        self.assertEqual(job["state"], "done", job)
        from anonymizer.addin import structure_of
        self.assertEqual(job["structure"]["slides"], structure_of(self.result(job["id"]), ".pptx")["slides"])
        self.assertGreaterEqual(job["structure"]["slides"], 1)

    def test_unsupported_and_empty_documents_are_refused_before_any_work(self):
        status, body = self.json("POST", "/api/addin/anonymize", b"abc", {"X-Filename": "a.txt"})
        self.assertEqual(status, 400)
        status, body = self.json("POST", "/api/addin/anonymize", b"", {"X-Filename": "a.docx"})
        self.assertEqual(status, 400)
        status, body = self.json("POST", "/api/addin/anonymize", b"not a zip", {"X-Filename": "a.docx"})
        self.assertEqual(status, 400)
        self.assertEqual(list((self.dir / "originals").glob("*.bin")), [])

    def test_word_tracked_changes_block_anonymizing_with_a_clear_code(self):
        import io
        import zipfile
        src = zipfile.ZipFile(io.BytesIO(docx_bytes([f"Директор {SECRET}."])))
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as z:
            for info in src.infolist():
                data = src.read(info.filename)
                if info.filename == "word/document.xml":
                    data = data.replace(b"<w:r>", b'<w:ins w:id="1" w:author="a" w:date="2026-01-01T00:00:00Z"><w:r>', 1) \
                               .replace(b"</w:r>", b"</w:r></w:ins>", 1)
                z.writestr(info, data)
        status, body = self.json("POST", "/api/addin/anonymize", out.getvalue(), {"X-Filename": "a.docx"})
        self.assertEqual((status, body.get("code")), (400, "tracked_changes"), body)
        self.assertGreaterEqual(body["count"], 1)
        self.assertEqual(self.app.uploads, {})                 # ничего не запущено и копия не сохранена
        self.assertEqual(list((self.dir / "originals").glob("*.bin")), [])
        # возврат при исправлениях не блокируется
        status, body = self.json("POST", "/api/addin/restore", out.getvalue(), {"X-Filename": "a.docx"})
        self.assertEqual(status, 200, body)

    def test_document_without_data_is_not_replaced(self):
        _, job = self.run_op("anonymize", "Пусто.docx", docx_bytes(["Просто текст без данных."]))
        self.assertEqual(job["state"], "done")
        self.assertEqual(job["files"][0]["total"], 0)

    def test_rerun_with_hidden_word_changes_the_result(self):
        text = "Согласовал Зарубин. Дальше работаем по плану."
        started, job = self.run_op("anonymize", "a.docx", docx_bytes([text]))
        status, again = self.json("POST", "/api/addin/rerun", json.dumps({"job": job["id"], "options": {}, "overrides": {"плану": "hide"}}).encode(),
                                  {"Content-Type": "application/json"})
        self.assertEqual(status, 200, again)
        self.assertEqual(again["backup"], started["backup"])
        for _ in range(300):
            status, new = self.json("GET", f"/api/addin/jobs/{again['job']}")
            if new["state"] != "running":
                break
            time.sleep(0.1)
        self.assertEqual(new["state"], "done", new)
        self.assertNotIn("плану", docx_text(self.result(again["job"])))

    def test_other_origin_cannot_read_results(self):
        _, job = self.run_op("anonymize", "a.docx", docx_bytes([f"Директор {SECRET}."]))
        status, _, _ = self.call("GET", f"/api/addin/result/{job['id']}", headers={"Origin": "https://evil.example"})
        self.assertEqual(status, 403)


class OriginalsTests(unittest.TestCase):
    def test_purge_keeps_newest_and_drops_old(self):
        env = Env()
        try:
            store = Originals(env.base / "o", env.key)
            ids = [store.save("anonymize", f"{i}.docx", b"x" * 10) for i in range(addin.KEEP_BACKUPS + 3)]
            store.purge()
            self.assertEqual(len(store.listing()), addin.KEEP_BACKUPS)
            self.assertIsNone(store.read("../etc"))
            self.assertIsNone(store.meta(ids[0]) if False else store.read("zzzzzzzz"))
        finally:
            env.close()

    def test_wrong_key_cannot_read_copy(self):
        env = Env()
        try:
            store = Originals(env.base / "o", env.key)
            backup = store.save("anonymize", "a.docx", b"secret")
            self.assertIsNone(Originals(env.base / "o", b"k" * 32).read(backup))
            self.assertEqual(store.read(backup), b"secret")
        finally:
            env.close()


class StructureTests(unittest.TestCase):
    def test_structure_of_bytes_and_paths_match(self):
        data = xlsx_bytes({"A1": 1}, title="Лист1", extra_sheets={"Итоги": {"A1": 2}})
        self.assertEqual([s["name"] for s in structure_of(data, ".xlsx")["sheets"]], ["Лист1", "Итоги"])

    def test_verify_notes_extra_words_without_failing(self):
        env = Env()
        try:
            a, b = env.base / "a.docx", env.base / "b.docx"
            a.write_bytes(docx_bytes(["один два"]))
            b.write_bytes(docx_bytes(["один два три"]))
            verdict = verify_document(a, b)
            self.assertTrue(verdict["ok"])
            self.assertTrue(verdict["notes"])
        finally:
            env.close()


if __name__ == "__main__":
    unittest.main()
