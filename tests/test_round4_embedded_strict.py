"""Вложения, которые нельзя обезличить (книга данных диаграммы .xlsb, объекты OLE), и режим «Скрывать и сомнительные слова».

Все названия, фамилии и содержимое вложений здесь выдуманы.
"""
import io
import re
import unittest
import zipfile
from pathlib import Path

import fitz
from docx import Document
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

from .helpers import Env, content_units, docx_bytes, pptx_bytes

WEB = Path(__file__).resolve().parent.parent / "anonymizer" / "web" / "index.html"

NS = ('xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
      'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
SECRET = "Зарубин Тихон Аркадьевич, ООО «Кедровый Лог»".encode("utf-8")
PAYLOAD = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1 fake ole " + SECRET


def png() -> bytes:
    return fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 4, 4), False).tobytes("png")


def edit_zip(data: bytes, change=None, add: dict | None = None) -> bytes:
    """`change(имя, байты)` возвращает новые байты, None (убрать часть) или (новое имя, байты)."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            raw = zin.read(info.filename)
            result = change(info.filename, raw) if change else raw
            if result is None:
                continue
            name, raw = result if isinstance(result, tuple) else (info.filename, result)
            zout.writestr(name, raw)
        for name, raw in (add or {}).items():
            zout.writestr(name, raw)
    return out.getvalue()


def with_defaults(types: bytes, *pairs) -> bytes:
    text = types.decode()
    for ext, mime in pairs:
        if f'Extension="{ext}"' not in text:
            text = text.replace("<Default ", f'<Default Extension="{ext}" ContentType="{mime}"/><Default ', 1)
    return text.encode()


def add_rels(rels: bytes, *items) -> bytes:
    extra = "".join(f'<Relationship Id="{rid}" Type="{REL}/{kind}" Target="{target}"/>' for rid, kind, target in items)
    return rels.decode().replace("</Relationships>", extra + "</Relationships>").encode()


def ole_frame(with_picture: bool = True) -> str:
    picture = ('<p:pic><p:nvPicPr><p:cNvPr id="0" name=""/><p:cNvPicPr/><p:nvPr/></p:nvPicPr>'
               '<p:blipFill><a:blip r:embed="rIdImg50"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
               '<p:spPr><a:xfrm><a:off x="914400" y="3657600"/><a:ext cx="1828800" cy="914400"/></a:xfrm>'
               '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>') if with_picture else ""
    obj = '<p:oleObj name="Лист" r:id="rIdOle50" imgW="1000" imgH="500" progId="Excel.Sheet.12"><p:embed/>'
    return (f'<p:graphicFrame {NS}><p:nvGraphicFramePr><p:cNvPr id="50" name="Объект 50"/>'
            '<p:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></p:cNvGraphicFramePr><p:nvPr/></p:nvGraphicFramePr>'
            '<p:xfrm><a:off x="914400" y="3657600"/><a:ext cx="1828800" cy="914400"/></p:xfrm>'
            '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/presentationml/2006/ole">'
            '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
            f'<mc:Choice xmlns:v="urn:schemas-microsoft-com:vml" Requires="v">{obj}</p:oleObj></mc:Choice>'
            f'<mc:Fallback>{obj}{picture}</p:oleObj></mc:Fallback></mc:AlternateContent>'
            '</a:graphicData></a:graphic></p:graphicFrame>')


def pptx_with_ole(with_picture: bool = True) -> bytes:
    base = pptx_bytes(["Итоги квартала: поставки выросли", "Ответственный: Зарубин Тихон Аркадьевич"], title="Отчёт")

    def change(name, raw):
        if name == "ppt/slides/slide1.xml":
            return raw.decode().replace("</p:spTree>", ole_frame(with_picture) + "</p:spTree>").encode()
        if name == "ppt/slides/_rels/slide1.xml.rels":
            return add_rels(raw, ("rIdOle50", "oleObject", "../embeddings/oleObject1.bin"),
                            ("rIdImg50", "image", "../media/image50.png"))
        if name == "[Content_Types].xml":
            return with_defaults(raw, ("bin", "application/vnd.openxmlformats-officedocument.oleObject"), ("png", "image/png"))
        return raw
    return edit_zip(base, change, {"ppt/embeddings/oleObject1.bin": PAYLOAD, "ppt/media/image50.png": png()})


def pptx_with_xlsb_chart() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Выручка по клиентам"
    data = CategoryChartData()
    data.categories = ["Зарубин Тихон Аркадьевич", "Волобуева Нина Степановна"]
    data.add_series("ООО «Кедровый Лог»", (10, 20))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 1500000, 6000000, 3500000, data)
    buffer = io.BytesIO()
    prs.save(buffer)

    def change(name, raw):
        if name.startswith("ppt/embeddings/") and name.endswith(".xlsx"):
            # Двоичная книга: программа её не разбирает, а данные в ней исходные.
            return name[:-5] + ".xlsb", b"PK fake xlsb " + SECRET
        if name.startswith("ppt/charts/_rels/"):
            return raw.replace(b".xlsx", b".xlsb")
        if name == "[Content_Types].xml":
            return with_defaults(raw, ("xlsb", "application/vnd.ms-excel.sheet.binary.macroEnabled.12"))
        return raw
    return edit_zip(buffer.getvalue(), change)


def docx_with_ole() -> bytes:
    base = docx_bytes(["Приложение: расчёт стоимости поставки", "Ответственный: Зарубин Тихон Аркадьевич"])
    obj = ('<w:p><w:r><w:object w:dxaOrig="1440" w:dyaOrig="720">'
           '<v:shape id="_x0000_i1025" type="#_x0000_t75" style="width:72pt;height:36pt" o:ole="">'
           '<v:imagedata r:id="rIdImg50" o:title=""/></v:shape>'
           '<o:OLEObject Type="Embed" ProgID="Excel.Sheet.12" ShapeID="_x0000_i1025" DrawAspect="Content" '
           'ObjectID="_1700000001" r:id="rIdOle50"/></w:object></w:r></w:p>')

    def change(name, raw):
        if name == "word/document.xml":
            text = raw.decode()
            for prefix, uri in (("v", "urn:schemas-microsoft-com:vml"), ("o", "urn:schemas-microsoft-com:office:office")):
                if f"xmlns:{prefix}=" not in text:
                    text = text.replace("<w:document ", f'<w:document xmlns:{prefix}="{uri}" ', 1)
            return text.replace("<w:sectPr", obj + "<w:sectPr", 1).encode()
        if name == "word/_rels/document.xml.rels":
            return add_rels(raw, ("rIdOle50", "oleObject", "embeddings/oleObject1.bin"), ("rIdImg50", "image", "media/image50.png"))
        if name == "[Content_Types].xml":
            return with_defaults(raw, ("bin", "application/vnd.openxmlformats-officedocument.oleObject"), ("png", "image/png"))
        return raw
    return edit_zip(base, change, {"word/embeddings/oleObject1.bin": PAYLOAD, "word/media/image50.png": png()})


def names(data: bytes) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return z.namelist()


def part(data: bytes, name: str) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return z.read(name).decode("utf-8")


def raw_zip(data: bytes) -> bytes:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return b"".join(z.read(n) for n in z.namelist())


def messages(job, level=None) -> list[str]:
    return [m["text"] for m in job.files[0].messages if level is None or m["level"] == level]


class Base(unittest.TestCase):
    def setUp(self):
        self.env = Env()

    def tearDown(self):
        self.env.close()

    def round_trip(self, name: str, data: bytes):
        job = self.env.anonymize((name, data))
        self.assertNotEqual(job.files[0].status, "error", job.files[0].messages)
        restored = self.env.restore(self.env.as_upload(job))
        self.assertEqual(restored.result["unknown"], {})
        tmp = Path(self.env.base) / name
        tmp.write_bytes(data)
        self.assertEqual(content_units(tmp), content_units(self.env.output(restored)))
        return job, restored


class ChartWorkbookTests(Base):
    def test_binary_chart_workbook_is_removed_and_chart_keeps_its_values(self):
        job, restored = self.round_trip("chart.pptx", pptx_with_xlsb_chart())
        out = self.env.output_bytes(job)
        self.assertFalse([n for n in names(out) if "/embeddings/" in n])
        self.assertNotIn(SECRET, raw_zip(out))
        chart = next(n for n in names(out) if re.match(r"ppt/charts/chart\d+\.xml$", n))
        chart_xml = part(out, chart)
        self.assertNotIn("externalData", chart_xml)
        self.assertNotIn("relationships/package", part(out, chart.replace("charts/", "charts/_rels/") + ".rels"))
        self.assertNotIn('Extension="xlsb"', part(out, "[Content_Types].xml"))
        # Подписи категорий и имя ряда в кэше диаграммы обезличены, числа на месте: диаграмма рисуется как раньше.
        for leaked in ("Зарубин", "Волобуева", "Кедровый"):
            self.assertNotIn(leaked, chart_xml)
        self.assertIn("<c:v>20</c:v>", chart_xml)
        shape = next(s for s in Presentation(io.BytesIO(out)).slides[0].shapes if s.has_chart)
        self.assertEqual(len(list(shape.chart.plots[0].categories)), 2)
        notice = " ".join(messages(job, "info"))
        self.assertIn("Удалены вложения", notice)
        self.assertIn("встроенный .xlsb", notice)
        self.assertIn("диаграммы (1)", notice.lower())
        self.assertEqual(job.files[0].status, "ok", job.files[0].messages)
        self.assertFalse(messages(job, "warn"))
        # Возврат: подписи диаграммы снова исходные, удалённая книга не возвращается.
        back = self.env.output(restored).read_bytes()
        self.assertIn("Зарубин Тихон Аркадьевич", part(back, chart))
        self.assertIn("Кедровый Лог", part(back, chart))

    def test_processable_chart_workbook_stays(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        data = CategoryChartData()
        data.categories = ["Север", "Юг"]
        data.add_series("Продажи", (1, 2))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, 4000000, 3000000, data)
        buffer = io.BytesIO()
        prs.save(buffer)
        job = self.env.anonymize(("chart.pptx", buffer.getvalue()))
        out = self.env.output_bytes(job)
        self.assertTrue([n for n in names(out) if n.startswith("ppt/embeddings/") and n.endswith(".xlsx")])
        self.assertTrue(any("externalData" in part(out, n) for n in names(out) if re.match(r"ppt/charts/chart\d+\.xml$", n)))
        self.assertNotIn("Удалены вложения", " ".join(messages(job)))


class OleObjectTests(Base):
    def test_slide_ole_object_becomes_its_picture(self):
        job, _ = self.round_trip("ole.pptx", pptx_with_ole())
        out = self.env.output_bytes(job)
        self.assertNotIn("ppt/embeddings/oleObject1.bin", names(out))
        self.assertIn("ppt/media/image50.png", names(out))
        self.assertNotIn(SECRET, raw_zip(out))
        slide = part(out, "ppt/slides/slide1.xml")
        self.assertNotIn("oleObj", slide)
        self.assertIn('r:embed="rIdImg50"', slide)
        rels = part(out, "ppt/slides/_rels/slide1.xml.rels")
        self.assertNotIn("oleObject", rels)
        self.assertIn("image50.png", rels)
        # Тип .bin по-прежнему нужен настройкам принтера из шаблона, поэтому его запись остаётся.
        self.assertIn("ppt/printerSettings/printerSettings1.bin", names(out))
        self.assertIn('Extension="bin"', part(out, "[Content_Types].xml"))
        shapes = {s.name: s for s in Presentation(io.BytesIO(out)).slides[0].shapes}
        self.assertEqual(shapes["Объект 50"].shape_type, 13)   # картинка
        self.assertEqual(shapes["Объект 50"].shape_id, 50)
        self.assertEqual(job.files[0].status, "ok", job.files[0].messages)
        self.assertIn("заменены их картинками", " ".join(messages(job, "info")))

    def test_ole_object_without_picture_is_kept_with_a_warning(self):
        job = self.env.anonymize(("ole.pptx", pptx_with_ole(with_picture=False)))
        out = self.env.output_bytes(job)
        self.assertIn("ppt/embeddings/oleObject1.bin", names(out))
        self.assertIn("oleObj", part(out, "ppt/slides/slide1.xml"))
        self.assertIn("oleObject1.bin", part(out, "ppt/slides/_rels/slide1.xml.rels"))
        self.assertTrue(any("объект OLE" in w for w in messages(job, "warn")), messages(job))
        self.assertNotIn("Удалены вложения", " ".join(messages(job)))
        Presentation(io.BytesIO(out))

    def test_word_ole_object_keeps_its_picture(self):
        job, _ = self.round_trip("ole.docx", docx_with_ole())
        out = self.env.output_bytes(job)
        self.assertNotIn("word/embeddings/oleObject1.bin", names(out))
        self.assertNotIn(SECRET, raw_zip(out))
        body = part(out, "word/document.xml")
        self.assertNotIn("OLEObject", body)
        self.assertIn('r:id="rIdImg50"', body)
        self.assertIn("<w:object", body)
        self.assertNotIn("oleObject", part(out, "word/_rels/document.xml.rels"))
        Document(io.BytesIO(out))
        self.assertEqual(job.files[0].status, "ok", job.files[0].messages)

    def test_shared_payload_still_used_elsewhere_is_kept(self):
        # Тот же объект нужен и части, где его заменить нечем: вложение остаётся, а о нём предупреждение.
        data = pptx_with_ole()

        def change(name, raw):
            if name == "ppt/slideLayouts/slideLayout1.xml":
                return raw.decode().replace("</p:spTree>", ole_frame(with_picture=False) + "</p:spTree>").encode()
            if name == "ppt/slideLayouts/_rels/slideLayout1.xml.rels":
                return add_rels(raw, ("rIdOle50", "oleObject", "../embeddings/oleObject1.bin"))
            return raw
        job = self.env.anonymize(("shared.pptx", edit_zip(data, change)))
        out = self.env.output_bytes(job)
        self.assertIn("ppt/embeddings/oleObject1.bin", names(out))
        self.assertIn('Extension="bin"', part(out, "[Content_Types].xml"))
        self.assertTrue(any("объект OLE" in w for w in messages(job, "warn")))
        Presentation(io.BytesIO(out))


class StrictModeTests(Base):
    TEXT = ("Протокол встречи.\nУчастники от подрядчика: ООО «Лазурь», АО «Северный Бетон», компания Orvenda. "
            "Со стороны заказчика отвечает Тумбасов.\n")

    def test_off_by_default_leaves_suggestions(self):
        job = self.env.anonymize(("m.txt", self.TEXT))
        out = self.env.output(job).read_text(encoding="utf-8")
        suggested = {s["text"] for s in job.result["suggestions"]}
        self.assertTrue({"Orvenda", "Тумбасов"} & suggested, suggested)
        for word in {"Orvenda", "Тумбасов"} & suggested:
            self.assertIn(word, out)
        self.assertEqual(job.result.get("strict_hidden"), 0)
        explicit = Env()
        try:
            off = explicit.anonymize(("m.txt", self.TEXT), options={"strict": False})
            self.assertEqual(explicit.output(off).read_text(encoding="utf-8"), out)
        finally:
            explicit.close()

    def test_strict_hides_every_suggestion_and_restores_exactly(self):
        job = self.env.anonymize(("m.txt", self.TEXT), options={"strict": True})
        out = self.env.output(job).read_text(encoding="utf-8")
        for word in ("Orvenda", "Тумбасов", "Лазурь", "Северный Бетон"):
            self.assertNotIn(word, out)
        self.assertEqual(job.result["suggestions"], [])
        self.assertGreaterEqual(job.result["strict_hidden"], 1)
        self.assertIn("Протокол встречи", out)
        restored = self.env.restore(self.env.as_upload(job))
        self.assertEqual(self.env.output(restored).read_text(encoding="utf-8"), self.TEXT)
        # Автоматически скрытое не записывается в «Всегда скрывать»: это не решение человека.
        self.assertEqual(self.env.vault.prefs["hide_terms"], [])

    def test_token_numbering_has_no_gaps_and_matches_manual_hide(self):
        strict = self.env.anonymize(("m.txt", self.TEXT), options={"strict": True})
        manual_env = Env()
        try:
            first = manual_env.anonymize(("m.txt", self.TEXT))
            hide = {s["text"]: "hide" for s in first.result["suggestions"]}
            manual = manual_env.service.rerun(first.id, {}, hide)
            self.assertEqual(self.env.output(strict).read_text(encoding="utf-8"),
                             manual_env.output(manual).read_text(encoding="utf-8"))
        finally:
            manual_env.close()

    def test_keep_terms_are_respected(self):
        env = Env(keep_terms=["Orvenda"])
        try:
            job = env.anonymize(("m.txt", self.TEXT), options={"strict": True})
            out = env.output(job).read_text(encoding="utf-8")
            self.assertIn("Orvenda", out)
            self.assertNotIn("Тумбасов", out)
        finally:
            env.close()

    def test_user_keep_decision_wins_over_strict(self):
        first = self.env.anonymize(("m.txt", self.TEXT), options={"strict": True})
        again = self.env.service.rerun(first.id, {}, {"Тумбасов": "keep"})
        out = self.env.output(again).read_text(encoding="utf-8")
        self.assertIn("Тумбасов", out)
        self.assertNotIn("Orvenda", out)

    def test_preference_is_used_when_option_is_absent(self):
        env = Env(strict=True)
        try:
            job = env.anonymize(("m.txt", self.TEXT))
            self.assertNotIn("Orvenda", env.output(job).read_text(encoding="utf-8"))
        finally:
            env.close()

    def test_page_has_the_checkbox_wired_like_numbers(self):
        html = WEB.read_text(encoding="utf-8")
        self.assertIn('id="opt-strict"', html)
        self.assertIn("Скрывать и сомнительные слова", html)
        self.assertIn('strict: $("#opt-strict").checked', html)
        self.assertIn('$("#opt-strict").checked = !!STATE.prefs.strict', html)
        self.assertIn('strict: $("#pref-strict").checked', html)
        # Флажок стоит в «Дополнительно» сразу за заменой чисел.
        self.assertLess(html.index('id="opt-numbers"'), html.index('id="opt-strict"'))
        self.assertLess(html.index('id="opt-strict"'), html.index('id="btn-anon"'))


if __name__ == "__main__":
    unittest.main()
