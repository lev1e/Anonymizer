"""Расширенный набор файлов для проверки: договор, протокол, письма, таблица сотрудников, презентация с диаграммой, PDF.

Все названия и люди вымышлены. Для каждого файла заданы:
  secrets — значения, которые в обезличенном файле встречаться не должны;
  keep    — обычные слова, которые заменять нельзя (ложные срабатывания).
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.comments import Comment
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


def inn10(prefix: str) -> str:
    digits = [int(c) for c in prefix]
    k = [2, 4, 10, 3, 5, 9, 4, 6, 8]
    return prefix + str(sum(a * b for a, b in zip(digits, k)) % 11 % 10)


def inn12(prefix: str) -> str:
    digits = [int(c) for c in prefix]
    k1 = [7, 2, 4, 10, 3, 5, 9, 4, 6, 8]
    k2 = [3, 7, 2, 4, 10, 3, 5, 9, 4, 6, 8]
    d1 = sum(a * b for a, b in zip(digits, k1)) % 11 % 10
    digits.append(d1)
    d2 = sum(a * b for a, b in zip(digits, k2)) % 11 % 10
    return prefix + str(d1) + str(d2)


def luhn(prefix: str) -> str:
    total = 0
    for index, ch in enumerate(reversed(prefix)):
        d = int(ch)
        if index % 2 == 0:
            d *= 2
            d = d - 9 if d > 9 else d
        total += d
    return prefix + str((10 - total % 10) % 10)


INN_ORG = inn10("770123456")
INN_PERSON = inn12("500100732259"[:10])
CARD = luhn("427600001234567")


@dataclass
class Case:
    name: str
    data: bytes
    secrets: list[str] = field(default_factory=list)
    keep: list[str] = field(default_factory=list)


def _docx(build) -> bytes:
    doc = Document()
    build(doc)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def contract() -> Case:
    def build(doc):
        doc.core_properties.author = "Смирнов Алексей Петрович"
        doc.core_properties.title = "Договор подряда с ООО «Вектор-Строй»"
        doc.add_heading("ДОГОВОР ПОДРЯДА № 14/2026", 1)
        doc.add_paragraph("г. Новосибирск                                                    15 марта 2026 года")
        doc.add_paragraph(
            "Общество с ограниченной ответственностью «Вектор-Строй», именуемое в дальнейшем «Подрядчик», в лице "
            "генерального директора Смирнова Алексея Петровича, действующего на основании Устава, с одной стороны, и "
            "Акционерное общество «Технопарк Сибирь», именуемое в дальнейшем «Заказчик», в лице заместителя директора "
            "Кузнецовой Марии Сергеевны, действующей на основании доверенности № 45, с другой стороны, заключили "
            "настоящий договор о нижеследующем.")
        doc.add_paragraph("1. Подрядчик обязуется выполнить работы по объекту «Каскад-2» по адресу: г. Новосибирск, "
                          "ул. Ленина, д. 15, оф. 304, а Заказчик обязуется принять и оплатить их.")
        doc.add_paragraph("2. Стоимость работ составляет 4 500 000 рублей, в том числе НДС 20%. Оплата производится "
                          "в течение 10 рабочих дней с даты подписания акта.")
        doc.add_paragraph("3. Контактные лица: со стороны Подрядчика — Смирнов А.П., тел. +7 (383) 210-45-67, "
                          "e-mail: a.smirnov@vector-stroy.ru; со стороны Заказчика — Кузнецова М.С., тел. 8 913 555 01 22, "
                          "e-mail: kuznetsova@technopark-sib.ru.")
        doc.add_paragraph("Реквизиты Подрядчика: ИНН " + INN_ORG + ", КПП 540101001, р/с 40702810900000012345 в "
                          "ПАО «Сибирьбанк», БИК 045004001.")
        table = doc.add_table(rows=3, cols=3)
        rows = [["Этап", "Срок", "Ответственный"], ["Проектирование", "до 30 апреля", "Сидорова Е.В."],
                ["Монтаж", "до 15 июня", "Петров Игорь Николаевич"]]
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
        doc.add_paragraph("Подрядчик: ____________ /Смирнов А.П./        Заказчик: ____________ /Кузнецова М.С./")
        doc.add_paragraph("Место подписания — офис Заказчика. Документ составлен в двух экземплярах, по одному для каждой стороны. "
                          "Белый цвет фасада согласован. Новый этап начинается в понедельник.")
    return Case("Договор подряда.docx", _docx(build),
                secrets=["Вектор-Строй", "Технопарк Сибирь", "Смирнов", "Кузнецов", "Сидорова", "Петров Игорь", "Каскад-2",
                         "Новосибирск", "Ленина", "a.smirnov@vector-stroy.ru", "kuznetsova@technopark-sib.ru", "210-45-67",
                         "555 01 22", INN_ORG, "40702810900000012345", "Сибирьбанк"],
                keep=["Подрядчик", "Заказчик", "Стоимость", "НДС", "Оплата", "Место подписания", "Белый цвет", "Новый этап",
                      "понедельник", "марта", "Проектирование", "Монтаж"])


def meeting() -> Case:
    def build(doc):
        doc.add_heading("Протокол встречи по проекту «Атлант»", 1)
        doc.add_paragraph("Дата: 12.02.2026. Место: переговорная 3, Zoom.")
        doc.add_paragraph("Присутствовали: Иванов И.И. (руководитель проекта), Петрова М.С. (аналитик), Козлов Дмитрий "
                          "(заказчик, ООО «Ромашка»), Anna Fischer (Siemens).")
        doc.add_paragraph("Решения:")
        for text in ["Иванов И.И. подготовит отчёт по KPI до 20 февраля.",
                     "Петровой М.С. передать выгрузку из ERP и SAP в формате Excel.",
                     "Козлов Д. согласует бюджет с финансовым директором, Морозовой Ольгой Викторовной.",
                     "Следующая встреча в Екатеринбурге, адрес уточнит Fischer."]:
            doc.add_paragraph(text, style="List Bullet")
        doc.add_paragraph("Риски: задержка поставки, курс валюты, перенос дедлайна на март. Мороз в январе задержал монтаж.")
    return Case("Протокол встречи.docx", _docx(build),
                secrets=["Иванов", "Петрова", "Петровой", "Козлов", "Ромашка", "Fischer", "Siemens", "Морозовой", "Атлант",
                         "Екатеринбург"],
                keep=["KPI", "ERP", "SAP", "Excel", "Zoom", "Решения", "Риски", "Мороз в январе", "март", "финансовым директором"])


def letter_ru() -> Case:
    text = ("Тема: Сроки поставки оборудования\n\n"
            "Добрый день, Ольга Николаевна!\n\n"
            "Направляю Вам уточнённый график поставки для завода в Красноярске. По договору с ООО «Сибирские Турбины» "
            "первая партия придёт до 5 мая. Прошу подтвердить получение письма и передать копию Виктору Андреевичу "
            "Лебедеву (v.lebedev@sibturbin.ru).\n\n"
            "Если возникнут вопросы, звоните по номеру +7 923 400-11-22 или в офис на ул. Мира, 7, Красноярск.\n\n"
            "С уважением,\nСергей Романов\nруководитель отдела продаж\nООО «ТрансЭнерго»\n")
    return Case("Письмо.txt", text.encode("utf-8"),
                secrets=["Ольга Николаевна", "Красноярск", "Сибирские Турбины", "Лебедев", "v.lebedev@sibturbin.ru",
                         "400-11-22", "Мира", "Романов", "ТрансЭнерго"],
                keep=["Тема", "Сроки поставки оборудования", "Добрый день", "руководитель отдела продаж", "С уважением",
                      "график поставки"])


def letter_en() -> Case:
    text = ("Subject: Quarterly review\n\nDear Mr. Johnson,\n\nThank you for meeting us in London last week. As discussed, "
            "Acme Corp. will share the updated forecast with Sarah Connor (sarah.connor@acme-corp.com, +44 20 7946 0958) "
            "by Friday. Please copy Michael Brown from Globex Ltd.\n\nBest regards,\nJohn Smith\nSenior Analyst, Northwind Traders\n")
    return Case("Letter.txt", text.encode("utf-8"),
                secrets=["Johnson", "London", "Acme", "Sarah Connor", "sarah.connor@acme-corp.com", "7946 0958", "Michael Brown",
                         "Globex", "John Smith", "Northwind"],
                keep=["Quarterly review", "Senior Analyst", "Best regards", "forecast", "Friday"])


def employees() -> Case:
    wb = Workbook()
    ws = wb.active
    ws.title = "Сотрудники"
    ws.append(["ФИО", "Должность", "Отдел", "Город", "Email", "Телефон", "Оклад"])
    people = [("Соколов Андрей Павлович", "Инженер", "Производство", "Омск", "a.sokolov@zavod-omsk.ru", "+7 913 111 22 33", 85000),
              ("Гришина Елена Игоревна", "Бухгалтер", "Финансы", "Томск", "e.grishina@zavod-omsk.ru", "+7 913 222 33 44", 72000),
              ("Абрамов Тимур Русланович", "Менеджер по продажам", "Продажи", "Москва", "t.abramov@zavod-omsk.ru", "+7 495 333 44 55", 98000),
              ("Орлова Наталья Викторовна", "Руководитель проектов", "Проектный офис", "Казань", "n.orlova@zavod-omsk.ru", "+7 843 444 55 66", 120000)]
    for person in people:
        ws.append(list(person))
    ws["I1"] = "Итого"
    ws["I2"] = "=SUM(G2:G5)"
    ws["A7"] = "Комментарий: Соколов А.П. переходит в Продажи с 1 апреля"
    ws["A2"].comment = Comment("Согласовано с Гришиной Е.И.", "Автор")
    hidden = wb.create_sheet("Скрытый")
    hidden["A1"] = "Резерв: Волков Пётр Сергеевич, ООО «Заря»"
    hidden.sheet_state = "hidden"
    wb.properties.creator = "Гришина Елена Игоревна"
    buffer = io.BytesIO()
    wb.save(buffer)
    return Case("Сотрудники.xlsx", buffer.getvalue(),
                secrets=["Соколов", "Гришина", "Абрамов", "Орлова", "Волков", "Заря", "a.sokolov@zavod-omsk.ru", "913 111 22 33",
                         "Омск", "Томск", "Казань", "Гришиной"],
                keep=["Инженер", "Бухгалтер", "Производство", "Финансы", "Продажи", "Оклад", "Итого", "Должность", "Отдел",
                      "Менеджер по продажам", "Руководитель проектов", "Проектный офис"])


def deck() -> Case:
    prs = Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = "Стратегия ООО «Горизонт» на 2027 год"
    title.placeholders[1].text = "Докладчик: Никитин Максим Олегович"
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Выручка по регионам"
    data = CategoryChartData()
    data.categories = ["Пермь", "Уфа", "Самара"]
    data.add_series("Выручка, млн руб.", (120, 95, 143))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(7), Inches(4), data)
    slide.notes_slide.notes_text_frame.text = "Не забыть упомянуть, что Никитин М.О. согласовал бюджет с Ковалёвой Ириной."
    third = prs.slides.add_slide(prs.slide_layouts[1])
    third.shapes.title.text = "Ключевые выводы"
    third.placeholders[1].text = "Рост выручки 12%\nЗапуск проекта «Северный ветер»\nКонтакт: n.nikitin@gorizont.ru"
    buffer = io.BytesIO()
    prs.save(buffer)
    return Case("Презентация.pptx", buffer.getvalue(),
                secrets=["Горизонт", "Никитин", "Пермь", "Уфа", "Самара", "Ковалёвой", "Северный ветер", "n.nikitin@gorizont.ru"],
                keep=["Выручка по регионам", "Ключевые выводы", "Рост выручки", "Докладчик", "Стратегия"])


PDF_FONT = Path("/System/Library/Fonts/Supplemental/Arial.ttf")


def pdf_case() -> Case | None:
    import fitz
    fonts = [PDF_FONT, Path("C:/Windows/Fonts/arial.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")]
    font = next((f for f in fonts if f.exists()), None)
    if font is None:
        return None
    doc = fitz.open()
    page = doc.new_page()
    lines = ["Справка о сотруднике", "Сотрудник: Ершов Павел Анатольевич", "Работодатель: ООО «Меридиан»",
             "Город: Нижний Новгород", "Телефон: +7 831 200 30 40", "Почта: p.ershov@meridian.ru"]
    for index, line in enumerate(lines):
        page.insert_text((72, 90 + index * 24), line, fontname="F0", fontfile=str(font), fontsize=12)
    data = doc.tobytes()
    doc.close()
    return Case("Справка.pdf", data, secrets=["Ершов", "Меридиан", "Нижний Новгород", "200 30 40", "p.ershov@meridian.ru"],
                keep=["Справка о сотруднике", "Работодатель", "Телефон"])


def all_cases() -> list[Case]:
    cases = [contract(), meeting(), letter_ru(), letter_en(), employees(), deck()]
    pdf = pdf_case()
    if pdf:
        cases.append(pdf)
    return cases
